#!/usr/bin/env python3
"""Where a school district's money comes from, by source, off the Census Bureau's FY2024 Annual
Survey of School System Finances -- Lunenburg's schools and Monty Tech, side by side.

    python3 scripts/census_revenue.py        # print both districts, every source, in dollars
    python3 scripts/census_revenue.py --items         # every local item: code, form line, name
    python3 scripts/census_revenue.py --verify-page   # the revenue page's item rows, recomputed

WHY. TJ, 16 October 2026: people in town say *"monty tech makes most of its money by grants and
fundraising"* -- *"we need to show that, and compare lunenburg's grants and how they compare."*
The district's own budget summary cannot answer it: grants and gifts are outside the operating
budget it sets out. The Census table is the one document here that counts EVERY revenue source
for both districts the same way, so it is the like-for-like comparison.

WHAT IS ASSERTED. The documentation (state-census/census-school24doc.docx) states that TOTALREV
is TFEDREV + TSTREV + TLOCREV, and lists the items each of those sums. All four identities are
checked for both districts, so a source missed here cannot hide in a remainder.

WHAT "GRANTS" MEANS HERE, and whose choice it is: OURS. Every federal dollar to a school district
is a grant, so `federal` is the narrow measure. `state_programs` adds the state's categorical
and competitive money (career and technical, school meals, all other, nonspecified), which is
where a state award such as a Clean Energy Center grant lands; it leaves out Chapter 70,
transportation reimbursement and the benefits the state pays on a district's behalf, none of
which anybody applies for. `donations` is the survey's own item, U50, private contributions.

WHAT "OTHER LOCAL" HOLDS. Local revenue less the town budgets (T02, D23) and private gifts (U50)
was one row on the page until TJ, 23 October 2026: *"that's a huge number and not to know what it
represents is hard... like what is A07 Tuition?! U97?!"* It is now its items, each named. The
NAME of each code is read at run time from the survey's own data dictionary (the TLOCREV entry of
the technical documentation, where each code sits beside its name), and its MEANING from the
FY2024 form F-33's instructions, Part I Section A, page 8 -- both census.gov's own files, landed
by scripts/fetch_census_f33.py. The form prints its codes in a column apart from its line labels,
so a code is never paired with a line by position on the form (CLAUDE.md 13b): the line is found
by the item's NAME, which the form and the dictionary both print. Each quoted instruction is
asserted to appear verbatim in the form's text, and an item with money in it and no entry in
LOCAL_ITEMS refuses the build.

SCHOOL MEALS ARE NOT GRANTS. Federal school-meal money (C25, Child Nutrition Act) and the state's
meal programs (C10) are paid per meal served, by formula. They once sat inside `grants` and were
a quarter of Monty Tech's "grants and donations"; they are now `meals`, a band of their own, and
`grants` is federal plus state programs LESS the meal items.
"""
import hashlib
import json
import os
import re
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
KEY = 'state-census/census-elsec24.xlsx'
DOC_KEY = 'state-census/census-school24doc.docx'
PATH = os.path.join(ROOT, 'sources', KEY)
CACHE = os.path.join(ROOT, 'build', 'cache')
FY = 2024
IDS = {'2507080': 'lunenburg', '2508050': 'monty'}
VERSION = 2
# PUPILS, for the page's per-pupil chart, from the SAME file and year as the dollars: V33, "Fall
# Membership" in the record layout (census-school24doc.txt line 477), "presented in whole
# amounts" (line 229) -- fall 2023 memberships from NCES's Common Core of Data (line 98). Each is
# held against DESE's own 1 October headcount for the same year, in dese-radar.csv, by
# PUPILS_CHECK: the Census may adjust an enrolment "inconsistent with the finances reported"
# (line 98), and a count that drifted would make every per-pupil figure wrong at once.
DESE_RADAR = os.path.join(ROOT, 'sources', 'data', 'dese-radar.csv')
DESE_LEA = {'lunenburg': '01620000', 'monty': '08320000'}
PUPILS_TOLERANCE = 0.03   # ours: within 3% of DESE's headcount, or the build stops

# The documentation's own lists, in its order.
FEDERAL = ('C14 C15 C19 C22 C23 C26 C27 B11 C20 C25 C36 B10 B12 B14 B13').split()
STATE = ('C01 C04 C05 C06 C07 C08 C09 C10 C11 C12 C13 C24 C35 C38 C39').split()
LOCAL = ('T02 T06 T09 T15 T40 T99 D11 D23 A07 A08 A09 A11 A13 A15 A20 A40 '
         'U11 U22 U30 U50 U97').split()
STATE_PROGRAMS = ('C04', 'C05', 'C06', 'C07', 'C08', 'C09', 'C10', 'C13', 'C24', 'C35', 'C39')
MEALS = ('C25', 'C10')   # federal and state school-meal programs, both the survey's own items
TOWNS = ('T02', 'D23')   # a parent town's appropriation; a regional district's member towns
# Local revenue that is neither the town budgets nor private gifts: the page's itemised rows.
OTHER_LOCAL = tuple(c for c in LOCAL if c not in TOWNS + ('U50',))

F33_KEY = 'state-census/census-f33_2024_508.pdf'
F33_TEXT = os.path.join(ROOT, 'sources', 'state-census', 'text', 'census-f33_2024_508.txt')
DOC_TEXT = os.path.join(ROOT, 'sources', 'state-census', 'text', 'census-school24doc.txt')

# EVERY LOCAL ITEM: code -> (the page's plain name, the form's line, the instruction, quoted).
#
# The Census NAME is not typed here: label() reads it off the data dictionary,
# census-school24doc.txt lines 558-599 (the TLOCREV entry: code, then its name). The FORM LINE
# is where the F-33 prints that same name -- page 1, text lines 41-58, and page 2, lines 92-95 --
# and form_line_names() checks every one of those pairings by name. The QUOTE is the form's own
# instruction for that line, page 8, "Part I - REVENUE, Section A - FROM LOCAL SOURCES", text
# lines 509-552 (A15's is Part IX, page 10, line 836). quote_ok() asserts each is verbatim.
LOCAL_ITEMS = {
    # Lines 1-5, text lines 511-517.
    'T06': ('Property taxes the district levies', 1,
            'Independent school districts should report their tax receipts here.'),
    'T09': ('Sales taxes the district levies', 2,
            'Lines 2 through 5 will be applicable to only a relatively small number of districts.'),
    'T15': ('Utility taxes the district levies', 3,
            'Lines 2 through 5 will be applicable to only a relatively small number of districts.'),
    'T40': ('Income taxes the district levies', 4,
            'Lines 2 through 5 will be applicable to only a relatively small number of districts.'),
    'T99': ('Other taxes the district levies', 5,
            'Taxes reported here should be those which the district has the power to levy.'),
    # Line 6, text lines 518-520.
    'T02': ('The parent town’s appropriation', 6,
            'Dependent school systems should report their tax receipts and any other amounts '
            'appropriated by their parent government on line 6'),
    # Line 7, text lines 521-526.
    'D23': ('Member towns’ assessments', 7,
            'Regional school systems should report assessments received from cities and towns here.'),
    # Line 8, text lines 527-532.
    'D11': ('Paid by other school systems', 8,
            'Report payments received from other school systems, both within and outside the '
            'state, for tuition, transportation, and other services (such as purchasing and data '
            'processing). Regional school systems should report assessments received from '
            'independent school districts here.'),
    # Lines 9-14, text lines 533-541.
    'A07': ('Tuition paid by students, parents or private sources', 9,
            'Report all student fees on these lines.'),
    'A08': ('Bus fees paid by students, parents or private sources', 10,
            'Report all student fees on these lines.'),
    'A11': ('Textbook sales and rentals', 11,
            'Report all student fees on these lines.'),
    'A09': ('School meal sales', 12,
            'Gross receipts from sale of school breakfasts, lunches, and milk (from students, '
            'teachers, adults, but not from state or federal funds) should be entered on line 12.'),
    'A13': ('Student activity receipts', 13,
            'Gross district activity receipts for those funds under control of the custodian of '
            'district funds should be included on line 13.'),
    'A20': ('Other sales and services', 14,
            'Revenues from other sales and services, including community services activities '
            'should be entered on line 14.'),
    # Part IX, line 1 -- a special processing item, not a Section A line. Text line 836.
    'A15': ('Student fees that cannot be split by kind', None,
            'Report here any item in your chart of accounts that cannot be crosswalked into the '
            'items identified in Part I, Lines 9 (tuition), 10 (transportation), 11 (textbook '
            'sales), 12 (school food service sales), or 13 (district activities).'),
    # Lines 15-20, text lines 542-552.
    'A40': ('Rent and use of school buildings and land', 15,
            'Include revenues received from allowing temporary possession or granting rights to '
            'the use of school district buildings, land, or other properties on line 15.'),
    'U11': ('Sale of property', 16,
            'Report amounts received from sale of real property, buildings (and improvements to '
            'them), land easements, rights-of-way, and other capital assets (buses, automobiles, '
            'etc.) on line 16.'),
    'U22': ('Interest earned', 17,
            'Include interest earnings from all funds held by the school system on line 17.'),
    'U30': ('Fines', 18,
            'Report revenues from penalties imposed for violations of law on line 18.'),
    'U50': ('Private gifts', 19,
            'Report gifts of cash or securities from private individuals or organizations on '
            'line 19.'),
    'U97': ('Refunds and other local revenue', 20,
            'Report receipts from refunds of prior year expenditures and other revenue from local '
            'sources not provided for elsewhere on line 20.'),
}


def fail(msg):
    raise SystemExit('census_revenue: ' + msg)


def _norm(t):
    """For comparing two printings of one name: case, spacing, the form's slashed zero, a
    trailing plural and the form's own NCES account numbers in brackets are not the name."""
    t = re.sub(r'\([^)]*\)?$', '', t.replace('Ø', '0')).lower()
    return ' '.join(re.sub(r'(?<=x)es$|s$', '', w) for w in re.findall(r'[a-z]+', t))


def _flat(t):
    return ' '.join(t.replace('\u2013', '-').split())


def labels():
    """{code: name} for every local item, off the data dictionary's TLOCREV entry -- the
    survey's own pairing of a code with its name."""
    if not os.path.exists(DOC_TEXT):
        fail('%s is not here; run scripts/fetch_census_f33.py' % os.path.relpath(DOC_TEXT, ROOT))
    lines = [l.strip() for l in open(DOC_TEXT, encoding='utf-8')]
    starts = [i for i, l in enumerate(lines) if l == 'TLOCREV' and i + 2 < len(lines)
              and lines[i + 2].startswith('(sum of')]
    if len(starts) != 1:
        fail('the data dictionary has %d TLOCREV entries with an item list, not one' % len(starts))
    out, i = {}, starts[0] + 3
    while i + 1 < len(lines) and re.fullmatch(r'[A-Z]\d\d', lines[i]):
        out[lines[i]] = lines[i + 1]
        i += 2
    if tuple(out) != tuple(LOCAL):
        fail('the data dictionary lists local items %s; this file sums %s' % (list(out), list(LOCAL)))
    return out


def form_line_names():
    """{line number: name} for Part I Section A of the F-33, off the form's own page 1-2 list."""
    if not os.path.exists(F33_TEXT):
        fail('%s is not here; run scripts/fetch_census_f33.py' % os.path.relpath(F33_TEXT, ROOT))
    out, on = {}, False
    for l in open(F33_TEXT, encoding='utf-8'):
        l = l.strip()
        if l.startswith('Section A') and 'FROM LOCAL SOURCES' in l:
            on = True
            continue
        if l.startswith('Section B'):
            on = False
        m = on and re.match(r'(\d+)\. (.+)$', l)
        if m:
            out[int(m.group(1))] = m.group(2)
    if sorted(out) != list(range(1, 21)):
        fail('the F-33 lists local lines %s, not 1 to 20' % sorted(out))
    return out


def local_items():
    """{code: dict(code, census, name, line, meaning)} for every local item, every pairing and
    every quotation checked against the two Census documents."""
    lab, form = labels(), form_line_names()
    text = _flat(open(F33_TEXT, encoding='utf-8').read())
    out = {}
    for code in LOCAL:
        if code not in LOCAL_ITEMS:
            fail('local item %s (%s) has no entry in LOCAL_ITEMS' % (code, lab[code]))
        name, line, quote = LOCAL_ITEMS[code]
        if line is not None and _norm(form[line]) != _norm(lab[code]):
            fail('F-33 line %d is %r; the dictionary names %s %r' % (line, form[line], code, lab[code]))
        if _flat(quote) not in text:
            fail('the F-33 instruction quoted for %s is not in the form, verbatim' % code)
        out[code] = dict(code=code, census=lab[code], name=name, line=line, meaning=quote)
    return out


def _rows():
    if not os.path.exists(PATH):
        fail('%s is not here; run sync_archive.py --pull.' % KEY)
    sha = hashlib.sha256(open(PATH, 'rb').read()).hexdigest()
    cp = os.path.join(CACHE, 'census-revenue-v%d-%s.json' % (VERSION, sha[:16]))
    if os.path.exists(cp):
        return json.load(open(cp, encoding='utf-8'))
    import openpyxl
    ws = openpyxl.load_workbook(PATH, read_only=True).worksheets[0]
    want = ('NAME', 'YRDATA', 'V33', 'TOTALREV', 'TFEDREV', 'TSTREV', 'TLOCREV') + tuple(FEDERAL + STATE + LOCAL)
    rows, head = {}, None
    for r in ws.iter_rows(values_only=True):
        if head is None:
            head = list(r)
            miss = [f for f in want if f not in head]
            if miss:
                fail('%s: no column %s' % (KEY, miss))
            continue
        x = dict(zip(head, r))
        if x['NCESID'] in IDS:
            rows[x['NCESID']] = {f: x[f] for f in want}
    os.makedirs(CACHE, exist_ok=True)
    with open(cp + '.tmp', 'w', encoding='utf-8') as fh:
        json.dump(rows, fh)
    os.replace(cp + '.tmp', cp)
    return rows


def revenue():
    """{'lunenburg': {...}, 'monty': {...}}, every figure in DOLLARS (the file's thousands times
    1,000), with the groupings the Monty Tech revenue page draws."""
    rows = _rows()
    named = local_items()
    out = {}
    for nid, who in IDS.items():
        x = rows.get(nid)
        if x is None or str(x['YRDATA']) != str(FY)[2:]:
            fail('%s: no FY%d row for NCES id %s' % (KEY, FY, nid))
        g = lambda items: sum(x[c] or 0 for c in items)
        for total, items in (('TFEDREV', FEDERAL), ('TSTREV', STATE), ('TLOCREV', LOCAL)):
            if x[total] != g(items):
                fail('%s %s: %s is %s, its items sum to %s' % (KEY, nid, total, x[total], g(items)))
        if x['TOTALREV'] != x['TFEDREV'] + x['TSTREV'] + x['TLOCREV']:
            fail('%s %s: TOTALREV is not federal + state + local' % (KEY, nid))
        k = lambda v: v * 1000
        d = dict(name=x['NAME'], nces=nid, fy=FY, pupils=x['V33'],
                 total=k(x['TOTALREV']), federal=k(x['TFEDREV']), state=k(x['TSTREV']),
                 local=k(x['TLOCREV']),
                 ch70=k(x['C01']), transport=k(x['C12']), on_behalf=k(x['C38']),
                 capital_aid=k(x['C11']),
                 state_programs=k(g(STATE_PROGRAMS)),
                 towns=k(g(TOWNS)), donations=k(x['U50']),
                 meals_federal=k(x['C25']), meals=k(g(MEALS)))
        if d['ch70'] + d['transport'] + d['on_behalf'] + d['capital_aid'] + d['state_programs'] != d['state']:
            fail('%s %s: the state groupings do not add to TSTREV' % (KEY, nid))
        d['other_local'] = d['local'] - d['towns'] - d['donations']
        d['other_local_items'] = {c: k(x[c] or 0) for c in OTHER_LOCAL}
        if sum(d['other_local_items'].values()) != d['other_local']:
            fail('%s %s: the other local items do not add to local less towns and gifts' % (KEY, nid))
        d['grants'] = d['federal'] + d['state_programs'] - d['meals']
        d['grants_and_donations'] = d['grants'] + d['donations']
        d['state_and_towns'] = d['ch70'] + d['towns']
        d['transport_and_pensions'] = d['transport'] + d['on_behalf']
        d['meals_and_local'] = d['meals'] + d['other_local']
        parts = (d['grants'], d['meals'], d['donations'], d['ch70'], d['transport'],
                 d['on_behalf'], d['capital_aid'], d['towns'], d['other_local'])
        if sum(parts) != d['total']:
            fail('%s %s: the page\'s nine parts do not add to TOTALREV' % (KEY, nid))
        unnamed = [c for c, v in d['other_local_items'].items() if v and c not in named]
        if unnamed:
            fail('%s %s: money in local items with no name: %s' % (KEY, nid, unnamed))
        out[who] = d
    for who, hc in dese_headcounts().items():
        p = out[who]['pupils']
        if not p or abs(p - hc) > PUPILS_TOLERANCE * hc:
            fail('%s: Census fall membership %s is not within %d%% of DESE\'s FY%d headcount %s'
                 % (who, p, round(PUPILS_TOLERANCE * 100), FY, hc))
        out[who]['dese_headcount'] = hc
    return out


def dese_headcounts():
    """{who: DESE's 1 October headcount for FY}, off dese-radar.csv's 'Student Headcount'."""
    import csv
    got = {}
    for r in csv.DictReader(open(DESE_RADAR, encoding='utf-8')):
        if r['measure'] == 'Student Headcount' and r['fy'] == str(FY):
            for who, lea in DESE_LEA.items():
                if r['lea'] == lea:
                    got[who] = int(float(r['value']))
    if set(got) != set(DESE_LEA):
        fail('dese-radar.csv has no FY%d headcount for %s' % (FY, sorted(set(DESE_LEA) - set(got))))
    return got


def verify_page():
    """Every itemised row on /monty-tech/monty-tech-revenue, recomputed by a different route: the
    workbook read afresh (no cache), by column NAME, and the published cells compared as text.
    And the rows together foot to the single "other local" figure they replaced."""
    import openpyxl
    page = json.load(open(os.path.join(ROOT, 'fy28', 'public', 'data', 'topics',
                                       'monty-tech-revenue.json'), encoding='utf-8'))
    ws = openpyxl.load_workbook(PATH, read_only=True).worksheets[0]
    it = ws.iter_rows(values_only=True)
    head = list(next(it))
    raw = {}
    for r in it:
        x = dict(zip(head, r))
        if x['NCESID'] in IDS:
            raw[IDS[x['NCESID']]] = x
    views = {v['key']: v for v in page['table']['views']}
    rows = [r for r in views['dollars']['rows'] if r.get('code')]
    bad = []
    usd = lambda v: '${:,}'.format(v)
    num = lambda t: int(t.strip('$').replace(',', ''))
    # PER PUPIL. Every row of the per-pupil view against its own dollars row over the file's
    # V33, read here afresh: footing may move a row by under a dollar and never more; the rows
    # add to the total row; the total is TOTALREV over V33; and the two rows the page's
    # sentences quote -- Chapter 70 (C01) and the town budgets (T02 + D23) -- are tied to their
    # Census items directly, not through the page's dollars.
    dol, pup = views['dollars']['rows'], views['pupil']['rows']
    if [r['label'] for r in dol] != [r['label'] for r in pup]:
        bad.append('the per-pupil view does not have the dollars view\'s rows')
    for col, who in ((0, 'monty'), (1, 'lunenburg')):
        x = raw[who]
        n = x['V33']
        for d, p_ in zip(dol, pup):
            exact = num(d['cells'][col]) / n
            if abs(num(p_['cells'][col]) - exact) >= 1:
                bad.append('%s %s a pupil: page %s, file %.2f' % (who, p_['label'], p_['cells'][col], exact))
        if sum(num(r['cells'][col]) for r in pup[:-1]) != num(pup[-1]['cells'][col]):
            bad.append('%s: the per-pupil rows do not add to the total a pupil' % who)
        if num(pup[-1]['cells'][col]) != round(x['TOTALREV'] * 1000 / n):
            bad.append('%s: total a pupil is not TOTALREV over V33' % who)
        by = {r['label']: r for r in dol}
        for lab, val in (('State Chapter 70 aid', x['C01']),
                         ('Town budgets, mostly property tax', (x['T02'] or 0) + (x['D23'] or 0))):
            if num(by[lab]['cells'][col]) != val * 1000:
                bad.append('%s %s: page %s, file %s' % (who, lab, by[lab]['cells'][col], usd(val * 1000)))
        print('census_revenue: %s, V33 fall membership %s; total %s a pupil, Chapter 70 %s, '
              'town budgets %s' % (who, n, pup[-1]['cells'][col],
                                   next(r for r in pup if r['label'] == 'State Chapter 70 aid')['cells'][col],
                                   next(r for r in pup if r['label'].startswith('Town budgets'))['cells'][col]))
    for col, who in ((0, 'monty'), (1, 'lunenburg')):
        x = raw[who]
        for r in rows:
            c = r['code'].split()[-1]
            if r['cells'][col] != usd((x[c] or 0) * 1000):
                bad.append('%s %s: page %s, file %s' % (who, c, r['cells'][col], usd((x[c] or 0) * 1000)))
        old = (x['TLOCREV'] - (x['T02'] or 0) - (x['D23'] or 0) - (x['U50'] or 0)) * 1000
        shown = sum(int(r['cells'][col].strip('$').replace(',', '')) for r in rows)
        if shown != old:
            bad.append('%s: the item rows sum to %s, the other local figure is %s' % (who, usd(shown), usd(old)))
        missing = [c for c in OTHER_LOCAL if x[c] and 'Census ' + c not in [r['code'] for r in rows]]
        if missing:
            bad.append('%s: items with money and no row: %s' % (who, missing))
    if bad:
        fail('the revenue page does not match the Census file:\n  ' + '\n  '.join(bad))
    print('census_revenue: %d itemised rows match the file, and foot to other local for both '
          'districts; every per-pupil row is within a dollar of the file' % len(rows))


if __name__ == '__main__':
    if sys.argv[1:] == ['--verify-page']:
        verify_page()
        sys.exit(0)
    if sys.argv[1:] == ['--items']:
        for c, it in local_items().items():
            print('%s  line %-4s %-52s %s' % (c, it['line'] or 'IX.1', it['census'], it['name']))
        sys.exit(0)
    R = revenue()
    for who, d in R.items():
        print('%s: Census V33 fall membership %s, DESE headcount %s' % (who, d['pupils'],
                                                                        d['dese_headcount']),
              file=sys.stderr)
    json.dump(R, sys.stdout, indent=1)
    print()
