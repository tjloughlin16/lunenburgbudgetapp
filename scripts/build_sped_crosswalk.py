#!/usr/bin/env python3
"""THE SPECIAL EDUCATION CROSSWALK: what DESE counts as special education, function code by
function code, beside what each of OUR three definitions counts, and where they differ.

    python3 scripts/build_sped_crosswalk.py           # write the CSV and the topic page
    python3 scripts/build_sped_crosswalk.py --check   # fail if either would change, or a join fails

    sources/data/sped-crosswalk.csv                     the permanent list
    fy28/public/data/topics/sped-definition.json        /special-education/sped-definition

TJ, 23 October 2026: "a permanent list somewhere accessible, to know EXACTLY which lines on
the budget and which function codes are considered special education and why, so it's always
100% consistent with DESE."

WHAT DESE SAYS SPECIAL EDUCATION IS. Not a function code. In DESE's chart a function is what
the money did (teaching, guidance, tuition); special education is a PROGRAM, a second
dimension. DESE's FY2025 crosswalk (`dese-eoyr-fy2025-function-crosswalk.xlsx`) lists, for
every function, the programs a district may report it under: some functions carry a
Special Education program column (2110, 2305, 2330, 3300, 9300 ...), the rest are
`Undistributed` only and so are never special education (administration, buildings,
insurance). The End of Year Report instructions define the program: "Special education
programs are for instruction of students with special needs, provided for under MGL Ch. 71B
and 603 CMR 28.00" and English learner programs are reported as regular day (printing
version, eoy25!B217 and B215). Read, never typed: both are quoted from the cells on every run.

WHAT THE DISTRICT REPORTED. The FY2023 and FY2024 End of Year Report WORKBOOKS
(`extract_eoyr.py`): Schedule 1's SPECIAL EDUCATION column, by function, for the school
committee (the general fund) and for the town (school choice and charter tuition, which the
state assesses on the town). Working copies passed on by records request, not the reports as
filed (rule 13a); every page that uses them says so. Schedule 4 (special education by
placement) and Schedule 7 (transportation) are read as cross-checks and their disagreements
with Schedule 1 are reported, not hidden.

OUR THREE DEFINITIONS. TJ, 23 October 2026: "if those match DESE then we need to include
them". Every item where a rule differed from DESE was settled against DESE's own documents
and written ONCE into `sources/data/sped-definition.csv` (model/sped_definition.py), which all
three rules now read: an item is counted only where DESE's documents count it, and left out
where they do not or where the district's two reports contradict each other. The `decision`
column names the item that decides each function; a function where a rule still differs from
DESE and no item decides it fails the build. Since then the three rules count the same total
in both years, and this page checks it:

  budget lines   model/sped.py is_sped() or is_tuition(), on the district budget book's own
                 FY2023 and FY2024 actual columns (`lps-budget-lines.csv`)
  journal        scripts/journal.py in_sped_total() -- THE special education account list,
                 with the DESE items
  sped costs     scripts/build_sped_costs.py group_of() -- segment 51, less two ELL accounts,
                 plus the special education transportation account, with the DESE items

The two ledger rules are applied to ONE ledger -- the Town's period-13 school general fund
report (`munis-school-ytd.csv`), spent plus encumbered, which is how DESE says to report a
year (printing version A184: encumbrances "should be recorded as expenditures") -- so a
difference between them is a difference of RULE, never of data.

THE GRAIN OF COMPARISON is DESE's function code and, where a town code has no DESE line,
Schedule 4's function group (Teaching (2300), Guidance (2700) ...), which is the grain at which
DESE itself reports special education. A town code that is not in DESE's FY2025 chart (2310,
2315, 2357, 2900) is carried as its own row, and joins its group by the first two digits --
OUR grouping, said in the row.

RULE 1. Every figure is an actual against an actual: the report's year against the ledger's
and the budget book's actual columns. Nothing here reads a budget, and nothing feeds the model.

RULE 7. Where DESE's figure and ours differ, the difference is measured. Why the district
allocated a mixed function the way it did is not in any document held, and it is registered
as a gap (`money-gaps.csv`) rather than explained.
"""
import argparse
import collections
import csv
import io
import json
import os
import re
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.append(os.path.join(ROOT, 'model'))   # after scripts/: model/ has its own conclusions.py
import conclusions as C  # noqa: E402
from conclusions import figure  # noqa: E402
import ladder as LD  # noqa: E402
import topic as TP  # noqa: E402
import topics as REG  # noqa: E402
from brief import step  # noqa: E402
import journal as J  # noqa: E402
import build_sped_costs as SC  # noqa: E402
import sped as SP  # noqa: E402
import sped_definition as SD  # noqa: E402  (THE DESE items, sources/data/sped-definition.csv)

DATA = os.path.join(ROOT, 'sources', 'data')
DESE_DIR = os.path.join(ROOT, 'sources', 'state-dese')
CROSSWALK_XLSX = os.path.join(DESE_DIR, 'dese-eoyr-fy2025-function-crosswalk.xlsx')
PRINTING_XLSX = os.path.join(DESE_DIR, 'dese-eoyr-fy2025-printing-version.xlsx')
S1 = os.path.join(DATA, 'eoyr-schedule1.csv')
S3 = os.path.join(DATA, 'eoyr-schedule3.csv')
S4 = os.path.join(DATA, 'eoyr-schedule4.csv')
S7 = os.path.join(DATA, 'eoyr-schedule7.csv')
LEDGER = os.path.join(DATA, 'munis-school-ytd.csv')
OUT_CSV = os.path.join(DATA, 'sped-crosswalk.csv')
OUT_DIR = os.path.join(ROOT, 'fy28', 'public', 'data', 'topics')
PAGE = 'sped-definition'
ME = os.path.basename(__file__)
YEARS = (2023, 2024)
Y0, Y1 = YEARS
TOL = 1.0          # dollars: the workbook prints whole dollars against a ledger in cents

# The instruction cells quoted as DESE's definition (printing version, sheet eoy25). The
# text is READ from these cells on every run and the check fails if a cell no longer says it.
DEF_CELLS = {'sped': ('B217', 'Special education programs are for instruction of students '
                              'with special needs'),
             'regular': ('B215', 'English language learner (ELL) programs'),
             'encumbrance': ('A184', 'the encumbrances should be recorded as expenditures'),
             'schedule4': ('A2151', 'Circuit Breaker expenditures must be included')}

# Town / district function codes that are NOT in DESE's FY2025 chart and that no DESE name
# maps. Each is a decision: a new one appearing in the ledger or the budget book on a counted
# line fails the build until it is declared here. Where the report files them is not stated
# in any document held (money-gaps.csv).
UNMAPPED_EXPECTED = {'2310', '2315', '2357', '2900'}
# Lines the FY2023-FY2024 report prints that DESE's FY2025 chart no longer holds.
REPORT_ONLY = {'9130'}    # Charter Transportation Tuition (Schedule 1 line 1796)

SPED_PUBLIC = 'Lunenburg Budget Project'
COLS = ['level', 'dese_function', 'dese_function_name', 'in_dese_chart', 'dese_group', 'item',
        'dese_program_coded', 'dese_status',
        'dese_sped_fy2023', 'dese_sped_fy2024', 'dese_all_programs_fy2023',
        'dese_all_programs_fy2024', 'dese_circuit_breaker_fy2023', 'dese_circuit_breaker_fy2024',
        'budget_lines_rule', 'budget_lines_fy2023', 'budget_lines_fy2024',
        'journal_rule', 'journal_fy2023', 'journal_fy2024',
        'sped_costs_rule', 'sped_costs_fy2023', 'sped_costs_fy2024',
        'counted_by_budget_lines', 'counted_by_journal', 'counted_by_sped_costs',
        'agrees_with_dese', 'why', 'evidence', 'decision']
DEFS = ('budget_lines', 'journal', 'sped_costs')
DEF_LABEL = {'budget_lines': 'budget lines (model/sped.py)', 'journal': 'journal (journal.py)',
             'sped_costs': 'sped costs (build_sped_costs.py)'}

GROUP_NAME = {}   # filled from Schedule 4's own labels


def fail(msg):
    raise SystemExit('build_sped_crosswalk: ' + msg)


def money(v):
    return float(v) if v not in (None, '') else 0.0


def r0(x):
    """A CSV amount: whole dollars as the report prints them, cents kept for the ledger."""
    return '%.2f' % x if abs(x - round(x)) > 0.004 else '%d' % round(x)


# ------------------------------------------------------------------------------ DESE

def dese_chart():
    """{code: (name, cell)} from the crosswalk's `codes` sheet, and {code: (programs, first
    cell)} from its Schedule 1/3/20 crosswalk."""
    wb = openpyxl.load_workbook(CROSSWALK_XLSX, read_only=True, data_only=True)
    codes = {}
    for i, r in enumerate(wb['codes'].iter_rows(min_row=1, values_only=True), 1):
        if i < 4 or not r[1]:
            continue
        codes[str(r[1]).strip()] = (str(r[2]).strip(), 'codes!B%d' % i)
    progs, first = collections.defaultdict(set), {}
    ws = wb['Sch1 Sch3 Sch20 Crosswalk']
    for i, r in enumerate(ws.iter_rows(min_row=1, values_only=True), 1):
        if i == 1 or not r[4]:
            continue
        f = str(r[4]).strip()
        progs[f].add(str(r[9]).strip())
        if str(r[9]).strip() == 'Special Education' and f not in first:
            first[f] = "'Sch1 Sch3 Sch20 Crosswalk'!J%d" % i
    wb.close()
    if len(codes) < 80:
        fail('DESE codes sheet read %d codes; the layout moved' % len(codes))
    for f in progs:
        if f not in codes:
            fail('the crosswalk names function %s, which its own codes sheet does not' % f)
    return codes, progs, first


def dese_definition():
    wb = openpyxl.load_workbook(PRINTING_XLSX, read_only=True, data_only=True)
    ws = wb['eoy25']
    want = {cell for cell, _ in DEF_CELLS.values()}
    got = {}
    for row in ws.iter_rows():
        for c in row:
            if getattr(c, 'coordinate', None) in want and isinstance(c.value, str):
                got[c.coordinate] = ' '.join(c.value.split())
    wb.close()
    out = {}
    for k, (cell, must) in DEF_CELLS.items():
        if must not in got.get(cell, ''):
            fail('printing version eoy25!%s no longer says %r' % (cell, must))
        out[k] = (cell, got[cell])
    return out


# THE DESE CELLS THE ITEMS REST ON (sources/data/sped-definition.csv, `dese_citation`). Read
# from the workbooks on every run; the build fails if a cell stops saying what the item cites.
ITEM_CELLS = [
    (PRINTING_XLSX, 'eoy25', 'D1584', 'EPIMS job codes 1211 through 1223, 1312 and 1320'),
    (PRINTING_XLSX, 'eoy25', 'D1810', 'Guidance Counselors, and School Adjustment Counselors (2710)'),
    (PRINTING_XLSX, 'eoy25', 'D1814', 'psychological evaluation (2800)'),
    (PRINTING_XLSX, 'eoy25', 'D958', 'must exceed the dollar limits for non-instructional equipment ($5,000)'),
    (PRINTING_XLSX, 'eoy25', 'D2170', 'DESE will record school choice tuition for special education by placement'),
    (PRINTING_XLSX, 'eoy25', 'D2172', 'DESE will record charter school tuition for special education by placement'),
    (PRINTING_XLSX, 'EOYR-EPIMScrosswalk', 'D28', 'Supervisor/Director of Pupil Personnel'),
    (PRINTING_XLSX, 'EOYR-EPIMScrosswalk', 'D29', 'Special Education Administrator'),
    (PRINTING_XLSX, 'EOYR-EPIMScrosswalk', 'D40', 'School Special Education Administrator'),
    (PRINTING_XLSX, 'EOYR-EPIMScrosswalk', 'D82', 'Guidance Counselor'),
    (PRINTING_XLSX, 'EOYR-EPIMScrosswalk', 'D84', 'School Adjustment Counselor \u2013 Non-Special Education'),
    (PRINTING_XLSX, 'EOYR-EPIMScrosswalk', 'D85', 'School Adjustment Counselor \u2013 Special Education'),
    (PRINTING_XLSX, 'EOYR-EPIMScrosswalk', 'D87', 'School Psychologist \u2013 Non-Special Education'),
    (PRINTING_XLSX, 'EOYR-EPIMScrosswalk', 'D88', 'School Psychologist \u2013 Special Education'),
    (PRINTING_XLSX, 'EOYR-EPIMScrosswalk', 'D89', 'School Social Worker \u2013 Non-Special Education'),
    (PRINTING_XLSX, 'EOYR-EPIMScrosswalk', 'D90', 'School Social Worker \u2013 Special Education'),
    (CROSSWALK_XLSX, 'codes', 'C46', 'Equipment (7300, 7400)'),
]


def item_cells():
    want = collections.defaultdict(dict)
    for path, sheet, cell, text in ITEM_CELLS:
        want[(path, sheet)][cell] = text
    for (path, sheet), cells in want.items():
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        got = {}
        for row in wb[sheet].iter_rows():
            for c in row:
                if getattr(c, 'coordinate', None) in cells and c.value is not None:
                    got[c.coordinate] = ' '.join(str(c.value).split())
        wb.close()
        for cell, text in cells.items():
            if text not in got.get(cell, ''):
                fail('%s %s!%s no longer says %r' % (os.path.basename(path), sheet, cell, text))
    return True


def schedule1_cells():
    """{(fy, cell): (function, column, amount)} -- every cell of Schedule 1, as extracted."""
    out = {}
    for r in csv.DictReader(open(S1, encoding='utf-8')):
        out[(int(r['fiscal_year'][2:]), r['cell'])] = (r['function_code'], r['column'],
                                                       money(r['amount']))
    return out


def cell_of(ref, s1c):
    """'eoy24!F126' -> its amount, failing if the extract does not hold the cell."""
    m = re.fullmatch(r'eoy(\d\d)!([A-Z]+\d+)', ref)
    if not m:
        fail('sped-definition.csv cites %r, which is not an eoyNN!CELL reference' % ref)
    k = (2000 + int(m.group(1)), m.group(2))
    if k not in s1c:
        fail('sped-definition.csv cites %s, which eoyr-schedule1.csv does not hold' % ref)
    return s1c[k][2]


def items_measured(led, bl):
    """Every claim the items' `why` makes about the report and the ledger, MEASURED here; the
    build fails if one stops holding. Returns the figures the page shows for each item."""
    item_cells()
    s1c = schedule1_cells()
    by = {it['id']: it for it in SD.ITEMS}
    for it in SD.ITEMS:
        for ref in it['report_cells']:
            cell_of(ref, s1c)
    acct = collections.defaultdict(float)          # (fy, account) -> spent
    for (fy, a, f, seg, desc, amt, *_r) in led:
        acct[(fy, a)] += amt
    func = lambda fy, code: sum(v for (y, a), v in acct.items() if y == fy and a.split('-')[3] == code)
    known = {a for (_, a) in acct}
    for it in SD.ITEMS:
        for a in it['accounts']:
            if a not in known:
                fail('sped-definition.csv %s names account %s, which the period-13 ledger does '
                     'not hold' % (it['id'], a))
    for x in SD.EXCLUDED_LINE_ITEMS:
        if not any((r['line_item'] or '').strip() == x for r in SP.rows()):
            fail('sped-definition.csv names budget line %r, which the book does not hold' % x)
    c = lambda ref: cell_of(ref, s1c)
    out = {}
    # coordinator: the report's special education salary in 2110 IS the account, both years
    (co,) = by['coordinator']['accounts']
    for fy in YEARS:
        if abs(c('eoy%d!F126' % (fy % 100)) - acct[(fy, co)]) > TOL:
            fail('coordinator: eoy%d!F126 is no longer the coordinator account' % (fy % 100))
    out['coordinator'] = {fy: acct[(fy, co)] for fy in YEARS}
    # spellings: the FY2023 report carries both lines the model used to miss
    sup = sum(a for (fy, cd, g, item, a, k) in bl if fy == Y0 and item == 'Spcl Ed General School Supplies')
    if abs(c('eoy23!F235') - sup) > TOL:
        fail('spellings: eoy23!F235 is no longer the special education school supplies line')
    rel = sum(v for (y, a), v in acct.items() if y == Y0 and a.startswith('0100-3-300-2310-51-1-06-2-53501')
              and a[-6:] in ('535011', '535012'))
    if abs(c('eoy23!F129') - rel) > TOL:
        fail('spellings: eoy23!F129 is no longer related services plus hospital tutoring')
    out['spellings'] = (sup, rel)
    # psychology: FY2023 all special education; FY2024 the same accounts in regular day, plus a
    # special education figure equal to FY2023's
    psy23 = func(Y0, '2800')
    if abs(c('eoy23!F274') + c('eoy23!F279') - psy23) > 2 * TOL:
        fail('psychology: FY2023 special education is no longer every 2800 account')
    sal24 = sum(v for (y, a), v in acct.items() if y == Y1 and a.split('-')[3] == '2800'
                and a.endswith('511023'))
    if abs(c('eoy24!E274') - sal24) > TOL:
        fail('psychology: eoy24!E274 is no longer the FY2024 psychologist salaries')
    if c('eoy24!F274') != c('eoy23!F274'):
        fail('psychology: the FY2024 special education figure no longer repeats FY2023\'s')
    if any(abs(v - c('eoy24!F274')) <= TOL for (y, a), v in acct.items() if y == Y1) \
            or abs(func(Y1, '2800') - c('eoy24!F274')) <= TOL:
        fail('psychology: the FY2024 special education figure now matches a FY2024 account')
    out['psychology'] = {Y0: psy23, Y1: func(Y1, '2800'), 'repeat': c('eoy24!F274')}
    # guidance: the same special education share both years, matching no account
    for ref in ('F258', 'F262'):
        if c('eoy23!' + ref) != c('eoy24!' + ref):
            fail('guidance: %s no longer repeats across the two reports' % ref)
    for fy in YEARS:
        if any(abs(v - c('eoy%d!F258' % (fy % 100))) <= TOL for (y, a), v in acct.items()
               if y == fy and a.split('-')[3] == '2710'):
            fail('guidance: a 2710 account now matches the report\'s special education share')
    out['guidance'] = {fy: c('eoy%d!F258' % (fy % 100)) + c('eoy%d!F262' % (fy % 100)) for fy in YEARS}
    # equipment: Schedule 1 gives line 1742 no special education cell
    if any(v[0] == '7300, 7400' and v[1] == 'SPECIAL EDUCATION' for v in s1c.values()):
        fail('equipment: the report now prints a special education cell for equipment')
    out['equipment'] = {fy: sum(v for (y, a), v in acct.items() if y == fy
                                and a.split('-')[3] == '7300' and a.split('-')[4] == '51')
                        for fy in YEARS}
    # choice and charter: the town's, both years
    out['choice-charter'] = {fy: c('eoy%d!F529' % (fy % 100)) + c('eoy%d!F530' % (fy % 100)) for fy in YEARS}
    if not all(out['choice-charter'].values()):
        fail('choice-charter: a year with no special education school choice or charter tuition')
    # the allocation: two substitute cells repeat across the reports too
    for ref in ('F173', 'F178'):
        if c('eoy23!' + ref) != c('eoy24!' + ref):
            fail('allocation: %s no longer repeats across the two reports' % ref)
    out['repeats'] = ['eoy24!F258', 'eoy24!F262', 'eoy24!F274', 'eoy24!F173', 'eoy24!F178']
    return out


def mapped_by_name(codes):
    """DESE codes whose NAME folds other codes in -- 'Equipment (7300, 7400)' -- as
    {folded code: DESE code}. DESE's own statement, read off the name."""
    out = {}
    for c, (name, _) in codes.items():
        m = re.search(r'\((\d{4}(?:,\s*\d{4})+)\)', name)
        if m:
            for x in re.findall(r'\d{4}', m.group(1)):
                if x != c:
                    out[x] = c
    return out


def schedule1():
    """{(fy, function): {'sc': sped, 'town': sped, 'all': all programs, 'cb': circuit breaker,
    'cells': [...]}} and the printed school committee special education total per year."""
    d = collections.defaultdict(lambda: dict(sc=0.0, town=0.0, all=0.0, cb=0.0, cells=[]))
    printed = {}
    for r in csv.DictReader(open(S1, encoding='utf-8')):
        fy = int(r['fiscal_year'][2:])
        f, col, sec = r['function_code'], r['column'], r['section']
        a = money(r['amount'])
        if sec.startswith('II. EXPENDITURES / A') and col == 'SPECIAL EDUCATION' \
                and r['description'].startswith('TOTAL EXPENDITURES BY SCHOOL COMMITTEE'):
            printed[fy] = (a, r['cell'])
        if not f or r['description'].startswith(('Sub-total', 'TOTAL')):
            continue
        x = d[(fy, f)]
        if sec.startswith('II. EXPENDITURES / A'):
            if col == 'SPECIAL EDUCATION':
                x['sc'] += a
                if a:
                    x['cells'].append('eoy%d!%s' % (fy % 100, r['cell']))
            elif col == 'TOTAL':
                x['all'] += a
        elif sec.startswith('II. EXPENDITURES / B'):
            if col == 'SPECIAL EDUCATION':
                x['town'] += a
                if a:
                    x['cells'].append('eoy%d!%s' % (fy % 100, r['cell']))
            elif col == 'TOTAL':
                x['all'] += a
        elif sec.startswith('C.2') and col.startswith('Revolving and Special Funds Circuit'):
            x['cb'] += a
            if a:
                x['cells'].append('eoy%d!%s (circuit breaker)' % (fy % 100, r['cell']))
    for fy in YEARS:
        if fy not in printed:
            fail('Schedule 1 FY%d: no printed school committee special education total' % fy)
        s = sum(v['sc'] for (y, _), v in d.items() if y == fy)
        if abs(s - printed[fy][0]) > TOL:
            fail('Schedule 1 FY%d: special education lines sum to %s, the report prints %s at %s'
                 % (fy, s, printed[fy][0], printed[fy][1]))
    return d, printed


def schedule3_total():
    """Schedule 3's line 3599 Special Education column: the instructional total DESE ties to."""
    out = {}
    for r in csv.DictReader(open(S3, encoding='utf-8')):
        if r['line'] == '3599' and r['column'] == 'Special Education':
            out[int(r['fiscal_year'][2:])] = (money(r['amount']), r['cell'])
    return out


def schedule4():
    """{(fy, line): {printed column: amount}} and {line: label} of Schedule 4."""
    d = collections.defaultdict(dict)
    for r in csv.DictReader(open(S4, encoding='utf-8')):
        d[(int(r['fiscal_year'][2:]), r['line'])][int(r['printed_column'])] = money(r['amount'])
        m = re.match(r'^(.*)\((\d{4})\)\s*$', r['description'])
        if m:
            GROUP_NAME.setdefault(m.group(2), m.group(1).strip())
    return d


def schedule7():
    d = collections.defaultdict(dict)
    for r in csv.DictReader(open(S7, encoding='utf-8')):
        d[(int(r['fiscal_year'][2:]), r['line'])][int(r['printed_column'])] = money(r['amount'])
    return d


# The Schedule 4 lines for each group of instructional functions, read off its own labels:
# `Supervisory (2100)` is line 3810 and so on. Derived from the CSV, asserted complete.
def s4_lines(s4):
    out = {}
    for r in csv.DictReader(open(S4, encoding='utf-8')):
        m = re.search(r'\((\d{4})\)\s*$', r['description'])
        if m:
            out[m.group(1)] = r['line']
    for g in ('2100', '2200', '2300', '2400', '2700', '2800', '9300', '9400'):
        if g not in out:
            fail('Schedule 4 has no line labelled (%s)' % g)
    return out


# ------------------------------------------------------------------------------ OURS

def group_of(code):
    """The grain DESE's Schedule 4 reports special education at: instructional functions by
    their first two digits (2110 -> 2100, `Supervisory (2100)`), everything else by itself."""
    return code[:2] + '00' if code.startswith('2') else code


def budget_lines():
    """[(fy, function code, function group, line item, amount, counted)] -- model/sped.py's
    rule on the budget book's actual columns. Function code is the group's leading four
    digits, as the book prints it."""
    out = []
    for r in SP.rows():
        g = (r['function_group'] or '').strip()
        m = re.match(r'^(\d{4})', g)
        if not m:
            continue
        counted = SP.is_sped(r) or SP.is_tuition(r)
        for fy in YEARS:
            a = money(r['fy%d_actual' % (fy % 100)])
            if a or counted:
                out.append((fy, m.group(1), g, r['line_item'], a, counted))
    return out


# A line the model does NOT count whose name still says special education in a spelling its
# rule does not read ('Spcl Ed'). Derived, and listed for TJ: the rule is model/sped.py's.
ABBREVIATED = re.compile(r'\bsp\w*\.? ?ed\b|\bsped\b', re.I)


def bl_reason(group, item, counted):
    """Why model/sped.py counts a budget line or does not, read off its own rule."""
    row = dict(function_group=group, line_item=item)
    if SP.is_tuition(row):
        return 'counted: tuition (9300, 9400)'
    if SP.NOT_SPED.search(item or ''):
        return 'not counted: English learner line (NOT_SPED), which DESE also reports as regular day'
    if counted:
        return ('counted: the group is special education (SPED_GROUPS)' if group in SP.SPED_GROUPS
                else 'counted: the line names special education')
    if ABBREVIATED.search(item or ''):
        return 'NOT counted, though the line names special education in a spelling the rule misses'
    return 'not counted'



def ledger():
    """[(fy, account, function, segment 5, description, amount, journal class, costs class)]
    for every school general fund expense account at period 13, FY2023 and FY2024."""
    A = J.accounts()
    out = []
    for r in csv.DictReader(open(LEDGER, encoding='utf-8')):
        if r['period'] != '13' or r['type'] != 'E' or r['report'] != 'gf-school' \
                or int(r['fiscal_year']) not in YEARS:
            continue
        a = A.get(r['account'])
        if a is None:
            fail('ledger account %s is not in the journal accounts file: the journal rule '
                 'cannot be applied to it' % r['account'])
        comp = J.component(a)
        jc = 'outside' if comp is not None and comp not in J.IN_TOTAL else (
            'counted' if comp is not None else 'no')
        g = SC.group_of(r['account'], r['description'])
        cc = 'counted' if g in SC.GROUPS else ('ell' if g == 'ell' else 'no')
        p = r['account'].split('-')
        out.append((int(r['fiscal_year']), r['account'], p[3], p[4], r['description'],
                    money(r['ytd_expended']) + money(r['encumbrances']), jc, comp or '', cc, g or ''))
    if not out:
        fail('no school general fund period-13 rows for FY%d-FY%d' % YEARS)
    return out


# ------------------------------------------------------------------------------ build

def build():
    codes, progs, first_sped = dese_chart()
    definition = dese_definition()
    folded = mapped_by_name(codes)
    s1, printed = schedule1()
    s3t = schedule3_total()
    s4 = schedule4()
    s7 = schedule7()
    s4l = s4_lines(s4)
    bl = budget_lines()
    led = ledger()

    def dese_code(code):
        """The DESE function a town/district code reports under, or None when not stated."""
        if code in codes:
            return code
        if ',' in code:   # Schedule 1 prints a folded line as '7100, 7200'
            hit = {folded.get(x.strip(), x.strip()) for x in code.split(',')} & set(codes)
            return hit.pop() if len(hit) == 1 else None
        return folded.get(code)

    # Every function Schedule 1 reports must be in DESE's chart (or folded into one).
    for (_, f) in s1:
        if dese_code(f) is None and f not in REPORT_ONLY:
            fail('Schedule 1 reports function %s, which DESE’s FY2025 chart does not hold' % f)
    # Re-key Schedule 1 on the DESE code each line reports under ('7100, 7200' -> 7200).
    merged = collections.defaultdict(lambda: dict(sc=0.0, town=0.0, all=0.0, cb=0.0, cells=[]))
    for (fy, f), v in s1.items():
        m = merged[(fy, dese_code(f) or f)]
        for k in ('sc', 'town', 'all', 'cb'):
            m[k] += v[k]
        m['cells'] += v['cells']
    s1 = dict(merged)

    # Every COUNTED line and account must land on a row (rule: a join that matches nothing
    # must fail). Unmapped town codes are allowed only when declared.
    counted_codes = ({c for (_, c, _, _, _, k) in bl if k}
                     | {f for (_, _, f, _, _, _, jc, _, cc, _) in led if jc != 'no' or cc != 'no'})
    stray = sorted(c for c in counted_codes if dese_code(c) is None and c not in UNMAPPED_EXPECTED)
    if stray:
        fail('special education lines or accounts in functions %s, which neither DESE’s chart '
             'nor UNMAPPED_EXPECTED holds -- a new code is a decision' % stray)
    unused = sorted(UNMAPPED_EXPECTED - {c for (_, c, *_r) in bl} - {x[2] for x in led})
    if unused:
        fail('UNMAPPED_EXPECTED declares %s, which nothing uses any more' % unused)

    # ---- amounts by (def, fy, row code) where the row code is the DESE code or the town code
    def row_code(c):
        return dese_code(c) or c
    ours = {d: collections.defaultdict(float) for d in DEFS}
    psych = collections.defaultdict(float)
    for (fy, c, g, item, a, k) in bl:
        if k:
            ours['budget_lines'][(fy, row_code(c))] += a
    for (fy, acct, f, seg, desc, a, jc, comp, cc, gg) in led:
        if jc == 'counted':
            ours['journal'][(fy, row_code(f))] += a
        if jc == 'outside' and comp == 'psychologists':
            psych[(fy, row_code(f))] += a
        if cc == 'counted':
            ours['sped_costs'][(fy, row_code(f))] += a

    # ---- THE DECISIONS: which sped-definition.csv item decides each function
    item_by = {it['id']: it for it in SD.ITEMS}
    item_of_code = {}
    for it in SD.ITEMS:
        for f in it['functions']:
            item_of_code.setdefault(f, it['id'])
        for a in it['accounts']:
            item_of_code.setdefault(row_code(SD.function_of(a)), it['id'])

    def decide(code, disagree, st, where):
        iid = item_of_code.get(code)
        if disagree and iid is None and code.startswith('2'):
            iid = 'allocation'       # the district's split of a mixed instructional function
        if disagree and iid is None:
            fail('%s: a rule differs from DESE and no sped-definition.csv item decides it -- '
                 'a new difference is a decision' % where)
        if iid is None:
            return 'none needed'
        it = item_by[iid]
        return '%s (sped-definition.csv: %s)' % (it['decision'], iid)

    measured = items_measured(led, bl)

    # ---- the rows
    rows = []
    all_codes = sorted(set(codes) | {row_code(c) for c in counted_codes}
                       | {row_code(c) for (_, c, *_r) in bl} | {row_code(x[2]) for x in led},
                       key=lambda c: (c[:1], c))

    def dese_amounts(code):
        v = {fy: s1.get((fy, code)) for fy in YEARS}
        sped = {fy: (v[fy]['sc'] + v[fy]['town']) if v[fy] else 0.0 for fy in YEARS}
        alls = {fy: v[fy]['all'] if v[fy] else 0.0 for fy in YEARS}
        cb = {fy: v[fy]['cb'] if v[fy] else 0.0 for fy in YEARS}
        cells = [c for fy in YEARS if v[fy] for c in v[fy]['cells']]
        return sped, alls, cb, cells

    def status(code, sped, alls):
        if code not in codes:
            return 'not a DESE code'
        if 'Special Education' not in progs.get(code, ()):
            return 'never (undistributed)'
        s, a = sped[Y1], alls[Y1]
        if a == 0 and s == 0:
            return 'nothing reported'
        if s == 0:
            return 'none'
        if abs(s - a) <= TOL:
            return 'all'
        return 'part'

    def agrees(d_sped, by_def):
        out = []
        for d in DEFS:
            ok = all(abs(by_def[d][fy] - d_sped[fy]) <= TOL for fy in YEARS)
            out.append('%s: %s' % (d, 'yes' if ok else 'no'))
        return '; '.join(out)

    def rule_text(code, which):
        """What a definition does in this function, derived from what it counted."""
        if which == 'budget_lines':
            grp = collections.OrderedDict()
            for (fy, c, g, item, a, k) in bl:
                if row_code(c) == code:
                    grp.setdefault(g, [0, 0])
                    grp[g][0 if k else 1] += 1
            if not grp:
                return 'no budget line in this function'
            parts = []
            for g, (yes, no) in grp.items():
                if yes and not no:
                    parts.append('%s: counted' % g)
                elif yes:
                    parts.append('%s: only lines naming special education (or tuition)' % g)
                else:
                    parts.append('%s: not counted' % g)
            return '; '.join(parts)
        key = 6 if which == 'journal' else 8
        seen = collections.Counter()
        for x in led:
            if row_code(x[2]) == code:
                seen[x[key]] += 1
        if not seen:
            return 'no general fund school account in this function'
        if which == 'journal':
            lab = {'counted': 'accounts on THE special education list', 'no': 'accounts not on it',
                   'outside': 'accounts listed, outside the total (sped-definition.csv)'}
        else:
            lab = {'counted': 'segment-51 accounts (and SPECIAL ED TRANSPORTATION)',
                   'no': 'accounts outside segment 51', 'ell': 'the ELL account(s), excluded'}
        return '; '.join('%d %s' % (seen[k], lab[k]) for k in lab if seen[k])

    totals = {d: {fy: 0.0 for fy in YEARS} for d in DEFS}
    dese_tot = {fy: 0.0 for fy in YEARS}
    groups = collections.OrderedDict()
    for code in all_codes:
        sped, alls, cb, cells = dese_amounts(code)
        by_def = {d: {fy: ours[d][(fy, code)] for fy in YEARS} for d in DEFS}
        for d in DEFS:
            for fy in YEARS:
                totals[d][fy] += by_def[d][fy]
        for fy in YEARS:
            dese_tot[fy] += sped[fy]
        in_chart = code in codes
        name = codes[code][0] if in_chart else ''
        if not in_chart:
            labels = sorted({g for (_, c, g, *_r) in bl if c == code})
            name = labels[0].split(' - ', 1)[-1] if labels else ''
        st = status(code, sped, alls)
        anything = any(sped.values()) or any(by_def[d][fy] for d in DEFS for fy in YEARS) \
            or any(psych[(fy, code)] for fy in YEARS)
        g = group_of(code)
        gr = groups.setdefault(g, dict(sped={fy: 0.0 for fy in YEARS},
                                       defs={d: {fy: 0.0 for fy in YEARS} for d in DEFS},
                                       codes=[]))
        for fy in YEARS:
            gr['sped'][fy] += sped[fy]
            for d in DEFS:
                gr['defs'][d][fy] += by_def[d][fy]
        if anything:
            gr['codes'].append(code)
        ag = agrees(sped, by_def) if (in_chart or anything) else ''
        if not in_chart:
            ag = 'not comparable here: DESE has no line for this code; compared in group %s' % g
        disagree = 'no' in ag or not in_chart
        why = []
        if not in_chart:
            why.append('Not a function in DESE’s FY2025 chart; which DESE line the district '
                       'reports it under is not stated in any document held. Grouped with %s by '
                       'its first two digits (our grouping).' % g)
        elif st == 'never (undistributed)':
            why.append('DESE’s crosswalk allows this function only as undistributed: it has '
                       'no program, so DESE never counts it as special education.')
        else:
            why.append('DESE reports this function by program; special education is the '
                       'program for students with special needs under MGL c.71B and 603 CMR '
                       '28.00 (%s), ELL is regular day (%s).'
                       % (definition['sped'][0], definition['regular'][0]))
            if st in ('all', 'part', 'none'):
                why.append('In FY%d the district reported %s of %s here as special education.'
                           % (Y1, C.usd(sped[Y1]), C.usd(alls[Y1])))
        town_codes = sorted(c for c in UNMAPPED_EXPECTED if group_of(c) == g and c != code)
        if in_chart and town_codes and anything:
            why.append('This group also holds the town code(s) %s, which DESE has no line for, '
                       'so compare at group %s rather than here.' % (', '.join(town_codes), g))
        if code in folded.values():
            why.append('DESE’s name folds in %s.' % ', '.join(sorted(k for k, v in folded.items()
                                                                  if v == code)))
        if any(psych[(fy, code)] for fy in YEARS):
            why.append('journal.py lists the psychologist accounts here (%s in FY%d) and keeps '
                       'them outside its total, by our choice.' % (C.usd(psych[(Y1, code)]), Y1))
        ev = []
        if in_chart:
            ev.append('dese-eoyr-fy2025-function-crosswalk.xlsx %s' % codes[code][1])
            if code in first_sped:
                ev.append('program Special Education allowed at %s' % first_sped[code])
        if cells:
            ev.append('eoyr-schedule1.csv cells ' + ', '.join(cells))
        if not (in_chart or anything):
            continue
        rows.append(dict(
            level='function', dese_function=code if in_chart else '', dese_function_name=name,
            in_dese_chart='yes' if in_chart else 'no', dese_group=g, item=code,
            dese_program_coded='yes' if 'Special Education' in progs.get(code, ()) else 'no',
            dese_status=st,
            **{'dese_sped_fy%d' % fy: r0(sped[fy]) for fy in YEARS},
            **{'dese_all_programs_fy%d' % fy: r0(alls[fy]) for fy in YEARS},
            **{'dese_circuit_breaker_fy%d' % fy: r0(cb[fy]) for fy in YEARS},
            budget_lines_rule=rule_text(code, 'budget_lines'),
            journal_rule=rule_text(code, 'journal'),
            sped_costs_rule=rule_text(code, 'sped_costs'),
            **{'%s_fy%d' % (d, fy): r0(by_def[d][fy]) for d in DEFS for fy in YEARS},
            counted_by_budget_lines='yes' if any(by_def['budget_lines'].values()) else 'no',
            counted_by_journal='yes' if any(by_def['journal'].values()) else 'no',
            counted_by_sped_costs='yes' if any(by_def['sped_costs'].values()) else 'no',
            agrees_with_dese=ag, why=' '.join(why), evidence='; '.join(ev),
            decision=decide(code, 'no' in ag, st, 'function %s' % code)))

        # ---- the rules within a function: every budget group and every ledger rule
        if not anything:
            continue
        sub = collections.OrderedDict()
        for (fy, c, gname, item, a, k) in bl:
            if row_code(c) != code:
                continue
            key = ('budget line group', '%s | %s' % (gname, bl_reason(gname, item, k)))
            s = sub.setdefault(key, dict(k=('budget_lines', k), amt={fy2: 0.0 for fy2 in YEARS},
                                         n=set()))
            s['amt'][fy] += a
            s['n'].add(item)
        for (fy, acct, f, seg, desc, a, jc, comp, cc, gg) in led:
            if row_code(f) != code:
                continue
            key = ('ledger account rule', 'segment %s | journal: %s | sped costs: %s'
                   % (seg, {'counted': 'counted (%s)' % comp, 'outside': 'listed, outside total '
                            '(%s)' % comp, 'no': 'not listed'}[jc],
                      {'counted': 'counted (%s)' % gg, 'ell': 'ELL, excluded', 'no': 'not counted'}[cc]))
            s = sub.setdefault(key, dict(k=('ledger', jc, cc), amt={fy2: 0.0 for fy2 in YEARS},
                                         n=set()))
            s['amt'][fy] += a
            s['n'].add(acct)
        for (lvl, item), s in sub.items():
            if not any(s['amt'].values()) and lvl == 'ledger account rule' \
                    and s['k'][1] == 'no' and s['k'][2] == 'no':
                continue
            r = {c: '' for c in COLS}
            r.update(level=lvl, dese_function=code if in_chart else '', dese_function_name=name,
                     in_dese_chart='yes' if in_chart else 'no', dese_group=g, item=item,
                     dese_program_coded='yes' if 'Special Education' in progs.get(code, ()) else 'no',
                     dese_status=st)
            if lvl == 'budget line group':
                k = s['k'][1]
                r.update(budget_lines_rule='counted' if k else 'not counted',
                         counted_by_budget_lines='yes' if k else 'no',
                         **{'budget_lines_fy%d' % fy: r0(s['amt'][fy]) for fy in YEARS},
                         evidence='lps-budget-lines.csv, line items: ' + '; '.join(sorted(s['n'])))
            else:
                _, jc, cc = s['k']
                r.update(journal_rule=jc, sped_costs_rule=cc,
                         counted_by_journal='yes' if jc == 'counted' else 'no',
                         counted_by_sped_costs='yes' if cc == 'counted' else 'no',
                         **{'journal_fy%d' % fy: r0(s['amt'][fy]) if jc == 'counted' else '0'
                            for fy in YEARS},
                         **{'sped_costs_fy%d' % fy: r0(s['amt'][fy]) if cc == 'counted' else '0'
                            for fy in YEARS},
                         evidence='munis-school-ytd.csv period 13, accounts: '
                                  + '; '.join(sorted(s['n'])))
                # the amount the rule touches, whoever counts it
                r['why'] = ('%s in FY%d and %s in FY%d on these accounts, spent plus encumbered.'
                            % (C.usd(s['amt'][Y0]), Y0, C.usd(s['amt'][Y1]), Y1))
            if lvl == 'budget line group':
                r['why'] = ('%s in FY%d and %s in FY%d in the book’s actual columns.'
                            % (C.usd(s['amt'][Y0]), Y0, C.usd(s['amt'][Y1]), Y1))
            r['agrees_with_dese'] = 'not comparable at this grain: DESE reports the function only'
            r['decision'] = rows[-1]['decision']
            rows.append(r)

    # ---- group rows: Schedule 4's grain, where the disagreements are measured
    data_codes = codes
    for g, gr in groups.items():
        if not gr['codes']:
            continue
        ag = agrees(gr['sped'], gr['defs'])
        s4line = s4l.get(g)
        ev = ['functions ' + ', '.join(gr['codes'])]
        if s4line:
            ev.append('Schedule 4 line %s: FY%d %s, FY%d %s (column 9, school committee, town '
                      'and circuit breaker)' % (s4line, Y0, C.usd(s4[(Y0, s4line)].get(9, 0)), Y1,
                                                 C.usd(s4[(Y1, s4line)].get(9, 0))))
        r = {c: '' for c in COLS}
        gname = GROUP_NAME.get(g) or data_codes.get(g, ('',))[0]
        r.update(level='group', dese_group=g, item=gname or g, dese_function_name=gname,
                 dese_function=g if g in data_codes else '',
                 **{'dese_sped_fy%d' % fy: r0(gr['sped'][fy]) for fy in YEARS},
                 **{'%s_fy%d' % (d, fy): r0(gr['defs'][d][fy]) for d in DEFS for fy in YEARS},
                 agrees_with_dese=ag, evidence='; '.join(ev),
                 why='Schedule 4’s grain: the functions in this group, DESE’s special education '
                     'column against each of our rules.',
                 decision='; '.join(sorted({decide(x, 'no' in ag, '', 'group %s' % g)
                                            for x in gr['codes']} - ({'none needed'} if 'no' in ag else set())))
                 or 'none needed')
        rows.append(r)

    # ---- the cross-checks inside the report
    checks = {}
    for fy in YEARS:
        instr1 = sum(v['sc'] + v['town'] for (y, f), v in s1.items() if y == fy and f.startswith('2'))
        if fy in s3t and abs(instr1 - s3t[fy][0]) > TOL:
            fail('FY%d: Schedule 1 instructional special education %s does not tie to Schedule 3 '
                 'line 3599 %s' % (fy, instr1, s3t[fy][0]))
        checks[fy] = dict(s1_instr=instr1, s4_instr=s4[(fy, '3860')].get(9, 0.0),
                          s1_trans=sum(v['sc'] for (y, f), v in s1.items() if y == fy and f == '3300'),
                          s7_trans=s7[(fy, '4160')].get(5, 0.0),
                          s4_screen=s4[(fy, '3860')].get(8, 0.0))
    missed = {fy: sum(a for (y, c, gname, item, a, k) in bl if y == fy and not k
                      and bl_reason(gname, item, k).startswith('NOT counted, though'))
              for fy in YEARS}
    data = dict(measured=measured, missed=missed, rows=rows, totals=totals, dese_tot=dese_tot, printed=printed, groups=groups,
                checks=checks, psych=psych, s1=s1, definition=definition, codes=codes,
                progs=progs, s4=s4, s4l=s4l)
    return data


def csv_text(rows):
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=COLS, lineterminator='\n')
    w.writeheader()
    for r in rows:
        w.writerow({c: r.get(c, '') for c in COLS})
    return buf.getvalue()


def report(d):
    """Every disagreement with DESE, in dollars by year, at the group grain."""
    out = ['DESE special education (school committee + town), FY%d %s, FY%d %s'
           % (Y0, C.usd(d['dese_tot'][Y0]), Y1, C.usd(d['dese_tot'][Y1]))]
    for dd in DEFS:
        out.append('  %-34s FY%d %s, FY%d %s' % (DEF_LABEL[dd], Y0, C.usd(d['totals'][dd][Y0]),
                                                   Y1, C.usd(d['totals'][dd][Y1])))
    out.append('Disagreements by group (ours minus DESE):')
    for g, gr in d['groups'].items():
        if not gr['codes']:
            continue
        for dd in DEFS:
            diffs = [gr['defs'][dd][fy] - gr['sped'][fy] for fy in YEARS]
            if any(abs(x) > TOL for x in diffs):
                out.append('  %s %-30s %-34s FY%d %s, FY%d %s' % (
                    g, (GROUP_NAME.get(g) or d['codes'].get(g, ('',))[0])[:30], DEF_LABEL[dd], Y0, C.usd(diffs[0]), Y1,
                    C.usd(diffs[1])))
    for fy in YEARS:
        c = d['checks'][fy]
        out.append('FY%d inside the report: Schedule 1 instructional %s vs Schedule 4 %s; '
                   'Schedule 1 transport %s vs Schedule 7 %s'
                   % (fy, C.usd(c['s1_instr']), C.usd(c['s4_instr']), C.usd(c['s1_trans']),
                      C.usd(c['s7_trans'])))
    return '\n'.join(out)


# ------------------------------------------------------------------------------ the page

def sources():
    return [
        dict(label='DESE — End of Year Financial Report, function and program crosswalk, '
                   'FY2025 (crosswalks.xlsx)',
             href='/docs/state-dese/dese-eoyr-fy2025-function-crosswalk.xlsx'),
        dict(label='DESE — End of Year Financial Report, FY2025, printing version, with the '
                   'instruction for each line (printing-version.xlsx)',
             href='/docs/state-dese/dese-eoyr-fy2025-printing-version.xlsx'),
        dict(label='Lunenburg Public Schools — End of Year Financial Report workbook, FY2024, '
                   'working copy dated 20 September 2024 (records request)',
             href='/docs/budget-workbooks/finance-committee/fy26-budget/department-presentations/'
                  'lunenburg-public-schools/school-budget-files-amanda-moore-foia-request-to-'
                  'school-dept/24eoy162-9-20-24.xlsx'),
        dict(label='Lunenburg Public Schools — End of Year Financial Report workbook, FY2023, '
                   'working copy (records request)',
             href='/docs/budget-workbooks/finance-committee/fy26-budget/department-presentations/'
                  'lunenburg-public-schools/school-budget-files-amanda-moore-foia-request-to-'
                  'school-dept/updated-th-23eoy162.xlsm'),
        dict(label='Lunenburg Budget Project — sped-crosswalk.csv, the whole list, every '
                   'function code, budget line group and ledger rule',
             href='/docs/data/sped-crosswalk.csv'),
        dict(label='Lunenburg Budget Project — munis-school-ytd.csv, the school general fund at '
                   'period 13', href='/docs/data/munis-school-ytd.csv'),
        dict(label='Lunenburg Budget Project — lps-budget-lines.csv, the district budget book, '
                   'line by line', href='/docs/data/lps-budget-lines.csv'),
    ]


def page(d):
    U = lambda v: figure(v, C.usd(v))
    F = lambda y: figure(y, C.fy(y), 'fiscal year')
    N = lambda n, unit: figure(n, C.num(n), unit)
    f, con = LD.fact, TP.conclusion
    g = d['groups']
    tot, dt = d['totals'], d['dese_tot']
    m = d['measured']
    # ONE DEFINITION: the three rules must now count the same, every year, or the page's first
    # sentence is false (sped-definition.csv is what they share).
    for fy in YEARS:
        if max(tot[x][fy] for x in DEFS) - min(tot[x][fy] for x in DEFS) > TOL:
            fail('FY%d: the three rules no longer agree (%s) -- the page says they do'
                 % (fy, {x: round(tot[x][fy]) for x in DEFS}))
    ours = {fy: tot['journal'][fy] for fy in YEARS}
    dsc = {fy: d['printed'][fy][0] for fy in YEARS}          # the school committee total, printed
    town = {fy: sum(d['s1'][(fy, c)]['town'] for c in ('9110', '9120')) for fy in YEARS}
    for fy in YEARS:
        if abs(dsc[fy] + town[fy] - dt[fy]) > TOL:
            fail('FY%d: DESE’s total is no longer the school committee plus the town' % fy)
    gap = {fy: dsc[fy] - ours[fy] for fy in YEARS}
    diff = lambda k, fy: g[k]['sped'][fy] - g[k]['defs']['journal'][fy]
    gui, psy, teach = diff('2700', Y1), diff('2800', Y1), diff('2300', Y1)
    if any(g['2700']['defs'][x][fy] or g['2800']['defs'][x][fy] for x in DEFS for fy in YEARS):
        fail('the page says no rule counts guidance or psychology; one now does')
    rest = gap[Y1] - gui - psy - teach
    if not (0 <= rest < min(gui, psy, teach)):
        fail('the page says guidance, psychology and teaching are most of the gap')
    agree_codes = ('3300', '9300', '9400')
    for code in agree_codes:
        for fy in YEARS:
            if abs(g[code]['defs']['journal'][fy] - g[code]['sped'][fy]) > TOL:
                fail('the page says %s agrees with DESE; FY%d does not' % (code, fy))
    ck = d['checks'][Y1]
    s4gap = ck['s1_instr'] - ck['s4_instr']
    nrep = len(m['repeats'])
    facts = [
        f('dsc', U(dsc[Y1]), 'special education, school committee, %s report to DESE' % C.fy(Y1),
          'published', {'y': F(Y1)}),
        f('ours', U(ours[Y1]), 'by our rules, all three, %s' % C.fy(Y1), 'published', {'y': F(Y1)}),
        f('gap', U(gap[Y1]), 'DESE counts and we do not, %s' % C.fy(Y1), 'published', {'y': F(Y1)}),
        f('dese', U(dt[Y1]), 'DESE’s whole %s total, the town’s tuition included' % C.fy(Y1),
          'published', {'y': F(Y1)}),
        f('town', U(town[Y1]), 'school choice and charter tuition, paid by the town', 'published'),
        f('gui', U(gui), 'of guidance DESE counts; we count none', 'published'),
        f('psy', U(psy), 'of psychology DESE counts; we count none', 'published'),
        f('teach', U(teach), 'more teaching and aides than our rules count', 'published'),
        f('rep', N(nrep, 'figures'), 'in the %s report that repeat %s’s exactly' % (C.fy(Y1), C.fy(Y0)),
          'published', {'y': F(Y1), 'z': F(Y0)}),
        f('tuit', U(sum(g[c]['sped'][Y1] for c in agree_codes)),
          'tuition and transport, %s: DESE and our rules alike' % C.fy(Y1), 'published',
          {'y': F(Y1)}),
        f('s4', U(s4gap), 'between two special education totals in one report', 'published'),
        f('alloc', None, 'which accounts make up DESE’s share of a mixed function', 'gap',
          closes='the district’s worksheet mapping each MUNIS account to an End of Year '
                 'Report line and program, FY2023 and FY2024'),
    ]
    xs = [k for k in g if g[k]['codes'] and (g[k]['sped'][Y1] or g[k]['defs']['journal'][Y1])]

    def label(k):
        n = GROUP_NAME.get(k) or d['codes'].get(k, ('',))[0] or k
        return n if k in n else '%s (%s)' % (n, k)
    chart = TP.series(
        'Special education by function, DESE and our rules', [label(k) for k in xs], [
            TP.track('DESE report', [g[k]['sped'][Y1] for k in xs], 'published', 'bar'),
            TP.track('Our rules', [g[k]['defs']['journal'][Y1] for k in xs], 'published', 'bar')],
        '$', kind='bar',
        note='%s, general fund. DESE’s special education column against our three rules, which '
             'now count the same, by function group.' % C.fy(Y1), note_figures={'y': F(Y1)})
    conclusions = [
        con('Our three rules now agree, at %s for %s; DESE’s report counts %s for the schools.'
            % (C.usd(ours[Y1]), C.fy(Y1), C.usd(dsc[Y1])),
            ['dsc', 'ours', 'gap'], {'a': U(ours[Y1]), 'y': F(Y1), 'b': U(dsc[Y1])},
            notes=[step('Each item where a rule differed from DESE is decided once, from DESE’s '
                        'documents, and every rule reads it: the table below.', 'measured'),
                   step('DESE’s whole total adds %s of school choice and charter tuition, which '
                        'the state charges to the town, not the schools.' % C.usd(town[Y1]),
                        'measured', {'a': U(town[Y1])}),
                   step('The report is the district’s working copy, not the filing; DESE’s '
                        'filed figures would supersede it.', 'gap', allow=(C.fy(Y0), C.fy(Y1)),
                        closes='the End of Year Financial Report as filed with DESE, '
                               '%s and %s' % (C.fy(Y0), C.fy(Y1)))]),
        con('Guidance %s, psychology %s and teaching %s are most of what DESE counts and we do not.'
            % (C.usd(gui), C.usd(psy), C.usd(teach)), ['gui', 'psy', 'teach', 'alloc'],
            {'a': U(gui), 'b': U(psy), 'c': U(teach)},
            notes=[step('Each is a share the district assigned inside a mixed function. DESE '
                        'assigns it person by person, by what each one does; the ledger records '
                        'accounts, not people.', 'measured'),
                   step('Which accounts make up those shares is not in the report or the ledger.',
                        'gap', closes='the district’s EPIMS job codes for its psychologists, '
                                      'counselors and teachers, and its account-to-report worksheet')]),
        con('The %s report repeats %s’s special education figures exactly in %s places.'
            % (C.fy(Y1), C.fy(Y0), C.num(nrep)), ['rep', 'psy', 'gui'],
            {'y': F(Y1), 'z': F(Y0), 'n': N(nrep, 'figures')},
            notes=[step('Guidance salaries and supplies, psychology and two substitute lines carry '
                        'the same figure both years; psychology’s matches no %s account.' % C.fy(Y1),
                        'measured', {'y': F(Y1)}),
                   step('Whether they were carried forward unchanged is not stated in the '
                        'working copy.', 'hypothesis')]),
        con('Out-of-district tuition and transportation agree with DESE to the dollar, both '
            'years.', ['tuit', 'town', 's4'],
            notes=[step('Inside the %s working copy, Schedule 4 totals %s less than Schedule 3; '
                        'DESE’s instructions tie the two.'
                        % (C.fy(Y1), C.usd(s4gap)), 'measured', {'y': F(Y1), 'a': U(s4gap)},
                        allow=('Schedule 4', 'Schedule 3'))]),
    ]
    # Table 1: every function with special education money, either side, both years.
    fr = [r for r in d['rows'] if r['level'] == 'function'
          and any(money(r[c]) for c in ('dese_sped_fy%d' % Y1, 'journal_fy%d' % Y1,
                                        'dese_sped_fy%d' % Y0, 'journal_fy%d' % Y0))]
    lab = lambda r: '%s %s' % (r['item'], r['dese_function_name'][:34])
    views = []
    for fy in (Y1, Y0):
        views.append(dict(
            key='fy%d' % fy, label=C.fy(fy), total_col=None,
            kinds=['published', 'ledger', 'ledger'],
            columns=['DESE report', 'Our rules', 'DESE less ours'],
            rows=[dict(label=lab(r), cells=[C.usd(money(r['dese_sped_fy%d' % fy])),
                                            C.usd(money(r['journal_fy%d' % fy])),
                                            C.usd(money(r['dese_sped_fy%d' % fy])
                                                  - money(r['journal_fy%d' % fy]))])
                  for r in fr],
            note='General fund, spent plus encumbered. A code with no DESE name is the town’s '
                 'own, compared in its group. The whole list, with the rule inside each '
                 'function, is the CSV.'))
    # Table 2: the items, as sources/data/sped-definition.csv decides them.
    shown = {
        'coordinator': lambda: 'Report’s special education salary is this account: %s (%s), %s (%s)'
        % (C.usd(m['coordinator'][Y0]), C.fy(Y0), C.usd(m['coordinator'][Y1]), C.fy(Y1)),
        'spellings': lambda: '%s: special education supplies (%s) and related services (%s) in its special education column'
        % (C.fy(Y0), C.usd(m['spellings'][0]), C.usd(m['spellings'][1])),
        'psychology': lambda: '%s: all of it, %s. %s: every account as regular day, plus %s that matches no account'
        % (C.fy(Y0), C.usd(m['psychology'][Y0]), C.fy(Y1), C.usd(m['psychology']['repeat'])),
        'guidance': lambda: '%s in both years, matching no account'
        % C.usd(m['guidance'][Y1]),
        'equipment': lambda: 'Filed as undistributed; %s and %s on the account'
        % (C.usd(m['equipment'][Y0]), C.usd(m['equipment'][Y1])),
        'choice-charter': lambda: 'Town’s spending: %s (%s), %s (%s)'
        % (C.usd(m['choice-charter'][Y0]), C.fy(Y0), C.usd(m['choice-charter'][Y1]), C.fy(Y1)),
        'allocation': lambda: 'Teaching and aides: %s more than our rules, %s' % (C.usd(teach), C.fy(Y1)),
    }
    items = dict(caption='Each item, as DESE’s documents decide it', views=[dict(
        key='items', label='Items', total_col=None,
        kinds=['ledger', 'published', 'published', 'published'],
        columns=['Our rules', 'What DESE says', 'DESE’s document', 'What the district’s report shows'],
        rows=[dict(label=it['item'],
                   cells=[it['decision'], it['dese_rule'], it['dese_citation'],
                          shown[it['id']]() if it['id'] in shown else '—'])
              for it in SD.ITEMS],
        note='sources/data/sped-definition.csv, read by all three rules. The report is the '
             'district’s working copy of its End of Year Financial Report.')])
    p = TP.page(
        PAGE,
        'Our rules now count what DESE’s documents count; DESE’s %s report still counts %s more.'
        % (C.fy(Y1), C.usd(gap[Y1])),
        facts, conclusions, chart=chart,
        table=dict(caption='Special education by function code, DESE beside our rules',
                   views=views),
        tables=[items],
        sources=sources(),
        finding_figures={'y': F(Y1), 'a': U(gap[Y1])})
    p['not_established'] = [
        'Which ledger accounts and budget lines make up DESE’s special education share of a '
        'mixed function: the report prints one figure per function and program.',
        'Which of the district’s psychologists and counselors are coded special education in '
        'EPIMS, which is what DESE’s rule turns on.',
        'Where the district reports its own codes 2310, 2315, 2357 and 2900, which are not in '
        'DESE’s chart.',
        'Whether the report as filed with DESE matches these working copies.',
    ]
    return p


def outputs():
    d = build()
    pages = {}
    if REG.TOPICS.get(PAGE, (None, None, None))[2] == ME:
        pages[os.path.join(OUT_DIR, PAGE + '.json')] = json.dumps(
            page(d), indent=1, ensure_ascii=False, sort_keys=True) + '\n'
    else:
        fail('topic %r is not registered to %s in scripts/topic_registry' % (PAGE, ME))
    pages[OUT_CSV] = csv_text(d['rows'])
    return pages, d


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--check', action='store_true')
    a = ap.parse_args()
    outs, d = outputs()
    print(report(d))
    if a.check:
        stale = [os.path.relpath(p, ROOT) for p, text in outs.items()
                 if not os.path.exists(p) or open(p, encoding='utf-8').read() != text]
        if stale:
            print('STALE %s' % ', '.join(stale), file=sys.stderr)
            return 1
        print('sped crosswalk: %d outputs current' % len(outs))
        return 0
    for p, text in outs.items():
        os.makedirs(os.path.dirname(p), exist_ok=True)
        tmp = p + '.tmp'
        with open(tmp, 'w', encoding='utf-8', newline='\n') as fh:
            fh.write(text)
        os.replace(tmp, p)
        print('wrote %s' % os.path.relpath(p, ROOT))
    return 0


if __name__ == '__main__':
    sys.exit(main())
