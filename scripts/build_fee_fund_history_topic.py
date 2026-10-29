#!/usr/bin/env python3
"""THE ATHLETIC FEE FUND, EVERY YEAR HELD: fund 1301 ("Chapter 658 School Athletics"), FY2011-FY2026.

    python3 scripts/build_fee_fund_history_topic.py           # write fy28/public/data/topics/fee-fund-history.json
    python3 scripts/build_fee_fund_history_topic.py --check   # fail if it would change

    fee-fund-history   /athletics/fee-fund-history   Fee fund history

On the Athletics shelf (scripts/topic_registry/athletics.py, through TOPIC_GENERATORS). Asked for
28 October 2026, beside the athletic director page (build_athletic_director_topic.py), modelled
on the after-school activities fund page (build_activities_fund_topics.py).

THE FEE FUND LEDGER PAGE IS FOLDED IN (28 October 2026, TJ approved). `fee-fund-ledger` was the
fund's CASH, FY2024-FY2026, off the cashbook, and opened on the same FY2025 entries this page
does. Its unique parts live here now: the cashbook's year-by-year cash (a second view of the
years table), the fund's lowest cash, and the two minutes it quoted. TOPIC_MERGED in
fy28/src/routes.ts sends the old address here.

YEAR-END FIGURES ON THREE BASES, never blended into one series:
  * FY2011-FY2023 and FY2024: the annual town reports (FY2024's Special Revenue Fund Balance
    Detail, read in sources/data/page-reads/fy2024-p23-special-revenue.csv);
  * FY2025: the FY2025 report's same schedule prints $131,239.09 -- the fund's CASH, which the
    cashbook closes FY2025 at to the cent, $20,991.20 above the fund balance the district's
    reconciliation opens FY2026 with (asserted below);
  * FY2026: the district's own reconciliation, school-funds-fy26.xlsx -- a sheet the district
    assembled (rule 13a: `stated`), though the chain /analysis/sitting-on-money proves from
    the period-13 ledgers reaches the same figure.

WHAT IS MEASURED, and refused if it stops being true:
  * FY2011-FY2023: the Town's special revenue schedule in each annual town report, as read
    and tied row by row in sources/data/special-revenue-read.csv (forward + receipts -
    disbursements = carried, `row_ties` yes). Each year's forward must equal the year before's
    carried, or the page refuses.
  * FY2014-FY2019: what the fee-fund side bought, itemised in the district's
    fy19-proposed-athletics-budget.pdf (sources/data/athletics-history.csv, side `revolving`).
    FY2018 and FY2019 there are budgets, not actuals, and are labelled so.
  * FY2023-FY2026: the school journal (journal.py). MONEY IN is the income accounts (USER
    CHARGES, high and middle school), net of refunds. MONEY OUT is the site's SPENT. Both tie,
    every year, to the cent, to the Town's period-13 special funds report (munis-school-ytd.csv).
  * THE MOVES: every journal entry with a line on fund 1301 and a line on the general fund
    (0100), read with both ends. Every one runs the same way, fee fund to town budget; the
    page refuses if one ever runs the other way unannounced.

WHAT IS NOT: what the fund paid for in FY2020-FY2022 (the annual reports print totals only),
whose pay the fund's payroll lines were (the journal records one line per account per pay run,
no person), and what any memo behind a move says. Registered in money-gaps.csv.

THE REDACTION (memory: redaction-principle-recipients-not-providers). Only an organisation
named in PAYEES gets a row. Every person the fund paid -- officials paid one by one, staff
reimbursed -- is inside "every other payee".
"""
import argparse
import csv
import json
import os
import re
import sys
from collections import defaultdict
from decimal import Decimal

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
import conclusions as C  # noqa: E402
from conclusions import figure  # noqa: E402
import ladder as LD  # noqa: E402
import topic as TP  # noqa: E402
import topics as REG  # noqa: E402
from brief import step  # noqa: E402
import journal as J  # noqa: E402

OUT_DIR = os.path.join(ROOT, 'fy28', 'public', 'data', 'topics')
ME = os.path.basename(__file__)
TID = 'fee-fund-history'
MINE = sorted(t for t, (_, _, gen) in REG.TOPICS.items() if gen == ME)
FUND = '1301'
GF = '0100'
Y = J.YEARS
ZERO = Decimal(0)
DATA = os.path.join(ROOT, 'sources', 'data')
SR = os.path.join(DATA, 'special-revenue-read.csv')
HIST = os.path.join(DATA, 'athletics-history.csv')
YTD = os.path.join(DATA, 'munis-school-ytd.csv')
CASH = os.path.join(DATA, 'fund-1301-cash-journal.csv')
PR24 = os.path.join(DATA, 'page-reads', 'fy2024-p23-special-revenue.csv')
AR25 = os.path.join(ROOT, 'sources', 'town-annual-reports', 'text', '4130-fy-2025-annual-town-report.txt')
AR25_LINE = '1301 CHAPT 658  $      131,239.09'
FUNDS26 = os.path.join(ROOT, 'sources', 'budget-workbooks', 'school-funds-fy26.xlsx')
MINUTES = os.path.join(ROOT, 'sources', 'meetings', 'text')
FY19_TXT = os.path.join(ROOT, 'sources', 'district-budget', 'text', 'fy19-proposed-athletics-budget.txt')
SR_NAME = 'Chapter 658 School Athletics'
PAY = '519999'
NAMED_MIN = Decimal(5000)
HOW_CLOSES = ('the fund’s account detail (MUNIS Report 1) for FY2020 to FY2022, which the '
              'annual reports summarise as one total a year')
WHO_CLOSES = ('payroll detail by employee for the fund’s salary account, FY2023 to FY2026, '
              'which the journal records as one line per pay run')
MEMO_CLOSES = ('the memos of 12 August 2024, 30 January, 2 May, 2 July and 20 August 2025 '
               'that the journal entries cite')

# WHO EACH ORGANISATION IS, declared one at a time (ours), by what the organisation IS.
PAYEES = {
    'bus': ('Bus companies', ['DEE BUS SERVICE INC', 'AA TRANSPORTATION IN', 'WILSON BUS LINES']),
    'arb': ('ArbiterSports (officials)', ['ARBITERSPORTS LLC']),
    'gear': ('Uniforms and equipment', ['PRIME TIME SPORTS', 'RIDDELL/ALL AMERICAN', 'US SPORTS',
                                        'VARSITY ATHLETIC', 'GIBNEY SALES & RECON', 'CROWN TROPHY',
                                        'PERFORMANCE HEALTH S', 'BC TECHNOLOGIES CO']),
    'video': ('Game video (Hudl)', ['AGILE SPORTS TECHNOL']),
    'site': ('Fields, venues and utilities', ['NASHOBA HILL CORP', 'SHREWSBURY LANDSCAPE',
                                              'LUNENBURG WATER DIST', 'GYMNASIUM FLOORS INC',
                                              'GEOSURFACES SOUTHEAS', 'WACHUSETT MOUNTAIN',
                                              'LUNENBURG SEWER DEPT']),
    'org': ('Other schools and leagues', ['LEOMINSTER PUBLIC SC', 'MASS STATE TRACK']),
}
GROUP_OF = {v: k for k, (_, vs) in PAYEES.items() for v in vs}
ORDER = ['pay', 'bus', 'arb', 'gear', 'video', 'site', 'org', 'rest']
LABEL = dict({k: n for k, (n, _) in PAYEES.items()}, pay='Pay', rest='Every other payee')


def fail(msg):
    raise SystemExit('build_fee_fund_history_topic: ' + msg)


def series(title, x, tracks, unit, note, **kw):
    """TP.series with the years and amounts a note names declared as its literals."""
    return TP.series(title, x, tracks, unit, note=note,
                     note_allow=tuple(re.findall(r'\d[\d,.]*\d|\d', note)), **kw)


def D(*texts):
    """The years and dates a text names, declared as literals (they are addresses, not figures)."""
    return tuple(re.findall(r'\d[\d,]*', ' '.join(texts)))


def usd2(v):
    v = Decimal(v)
    return ('-$%s' if v < 0 else '$%s') % format(abs(v), ',.2f')


def long_date(iso):
    y, m, d = (int(x) for x in iso.split('-'))
    return '%d %s %d' % (d, ('January February March April May June July August September October '
                             'November December').split()[m - 1], y)


# ------------------------------------------------------------------ the annual reports
def annual():
    rows = [r for r in csv.DictReader(open(SR, encoding='utf-8')) if r['fund'] == SR_NAME]
    out = {}
    for r in rows:
        y = int(r['fy'])
        if y in out:
            fail('two annual-report rows for FY%d' % y)
        if r['row_ties'] != 'yes':
            fail('FY%d: the annual report row does not tie' % y)
        f, rc, d, c = (Decimal(r[k]) for k in ('forward', 'receipts', 'disbursements', 'carried'))
        if f + rc - d != c:
            fail('FY%d: forward + receipts - disbursements is not carried' % y)
        out[y] = dict(forward=f, receipts=rc, disb=d, carried=c, page=r['page'], doc=r['document'])
    ys = sorted(out)
    if ys != list(range(ys[0], ys[-1] + 1)):
        fail('the annual-report years are not consecutive: %s' % ys)
    for a, b in zip(ys, ys[1:]):
        if out[b]['forward'] != out[a]['carried']:
            fail('FY%d opens at %s; FY%d closed at %s' % (b, out[b]['forward'], a, out[a]['carried']))
    return out


# ------------------------------------------------------------------ the FY19 document
def fy19():
    """fy -> {item: (amount, basis)} for the fee-fund side, and the raw line that names the
    Athletic Secretary (658), which the page quotes."""
    out = defaultdict(dict)
    for r in csv.DictReader(open(HIST, encoding='utf-8')):
        if r['side'] == 'revolving' and 'fy19-proposed' in r['source']:
            out[int(r['fy'])][r['item']] = (Decimal(r['amount']), r['basis'])
    raw = [l for l in open(FY19_TXT, encoding='utf-8') if l.startswith('Athletic Secretary (658)')]
    if len(raw) != 1:
        fail('the FY19 document no longer prints one Athletic Secretary (658) line')
    if any(l.startswith('Athletic Director') and '(658)' in l for l in open(FY19_TXT, encoding='utf-8')):
        fail('the FY19 document prints a fee-fund director line; the page says it prints none')
    for y in (2015, 2016, 2017):
        if out[y].get('Athletic Secretary', (None, None))[1] != 'actual':
            fail('FY%d: the FY19 document no longer gives the secretary as an actual' % y)
    return out, raw[0].strip()


# ------------------------------------------------------------------ the journal
def journal():
    A = J.accounts()
    L = J.lines()
    mine = [r for r in L if r['fund'] == FUND]
    if not mine:
        fail('the journal holds no line on fund %s' % FUND)
    jr_both = {r['journal'] for r in mine if r['source'] in ('GEN', 'GRV')}
    moves = []
    for jr in sorted(jr_both):
        rs = [r for r in L if r['journal'] == jr]
        funds = {r['fund'] for r in rs}
        if funds != {FUND, GF}:
            continue                         # a payroll accrual across many funds: not a move
        f_side = sum((r['amt'] for r in rs if r['fund'] == FUND), ZERO)
        g_side = sum((r['amt'] for r in rs if r['fund'] == GF), ZERO)
        if f_side + g_side != 0:
            # an accrual (both sides charged, the other end a liability outside the school
            # journal): not a move. A memo entry must never land here.
            if 'MEMO' in rs[0]['description'].upper():
                fail('%s cites a memo and does not move its money end to end' % jr)
            continue
        if any(not J.is_expense(A[r['account']]) for r in rs):
            fail('%s moves money on an account that is not an expense' % jr)
        moves.append(dict(journal=jr, fy=rs[0]['fy'], date=rs[0]['effective_date'],
                          source=rs[0]['source'], memo=rs[0]['description'], amount=-f_side,
                          to=sorted({(r['account_description'], r['amt']) for r in rs if r['fund'] == GF},
                                    key=lambda x: -x[1]),
                          frm=sorted({r['account_description'] for r in rs if r['fund'] == FUND})))
    moved_j = {m['journal'] for m in moves}
    if any(m['amount'] < 0 for m in moves if m['source'] == 'GEN'):
        fail('an entry moved money from the town budget onto the fee fund; the page says none did')

    t = defaultdict(Decimal)
    for r in mine:
        a, y = A[r['account']], r['fy']
        if J.is_revenue(a):
            if r['posts_to'] == 'actual':
                t[('in', y)] -= r['amt']
            continue
        if not J.is_expense(a):
            fail('%s is neither income nor expense' % r['account'])
        if not J.is_spent(r):
            continue
        if r['posts_to'] == 'encumbrance':
            t[('enc', y, r['account'])] += r['amt']
            if y != J.OPEN_FY:
                continue
        t[('out', y)] += r['amt']
        if r['journal'] in moved_j:
            t[('moved', y)] += r['amt']
            continue
        t[('booked', y)] += r['amt']
        if r['object'] == PAY:
            g = 'pay'
            t[('pay', y)] += r['amt']
        else:
            p = J.payee(r)
            g = GROUP_OF.get(p, 'rest') if isinstance(p, str) else 'rest'
            if g != 'rest':
                t[('payee', y, p)] += r['amt']
        t[('g', y, g)] += r['amt']
    for k in [k for k in t if k[0] == 'enc' and k[1] != J.OPEN_FY]:
        if t[k] != 0:
            fail('FY%d %s: orders do not net to nothing over a closed year' % (k[1], k[2]))
    for y in Y:
        if sum((t[('g', y, g)] for g in ORDER), ZERO) != t[('booked', y)]:
            fail('FY%d: the groups do not sum to what was booked' % y)
        if t[('booked', y)] + t[('moved', y)] != t[('out', y)]:
            fail('FY%d: booked plus moved is not spent' % y)
    # the pay runs: one regular slice beside the season lumps (read on the athletic director page)
    tie(t, mine, A)
    return t, moves


def tie(t, L, A):
    """THE TOWN'S OWN TOTALS: the period-13 special funds report prints fund 1301 account by
    account. Every year the journal's actual expense postings equal its `ytd_expended`, and its
    income postings what it prints received, to the cent. Fund 1301 carries no interest account
    on that report (1305 does), so there is nothing to set aside."""
    p = defaultdict(Decimal)
    for r in csv.DictReader(open(YTD, encoding='utf-8')):
        if r['fund'] != FUND or r['period'] != '13' or r['report'] != 'special-school':
            continue
        y = int(r['fiscal_year'])
        p[(y, r['type'])] += Decimal(r['ytd_expended'])
        p[(y, 'enc')] += Decimal(r['encumbrances'])
    for y in Y:
        je = sum((r['amt'] for r in L if r['fy'] == y and r['posts_to'] == 'actual'
                  and J.is_expense(A[r['account']])), ZERO)
        jr = sum((r['amt'] for r in L if r['fy'] == y and r['posts_to'] == 'actual'
                  and J.is_revenue(A[r['account']])), ZERO)
        if not p[(y, 'E')]:
            fail('FY%d: the period-13 report holds no fund %s' % (y, FUND))
        if je != p[(y, 'E')] or p[(y, 'enc')] != 0:
            fail('FY%d: journal spent %s; the report prints %s' % (y, je, p[(y, 'E')]))
        if jr != p[(y, 'R')]:
            fail('FY%d: journal received %s; the report prints %s' % (y, jr, p[(y, 'R')]))
        if -jr != t[('in', y)]:
            fail('FY%d: money in is not the report’s receipts' % y)
        t[('town_out', y)] = p[(y, 'E')]


# ------------------------------------------------------------------ the cashbook (from fee-fund-ledger)
def cashbook():
    """fy -> opening, money in, paid out, journal entries, closing, and the lowest running cash
    (our arithmetic, postings in effective-date order) -- off our redacted extract of the
    Town's MUNIS journal export for fund 1301. Each year's close must be the next one's
    printed opening (SOY BAL), or the page refuses."""
    rows = list(csv.DictReader(open(CASH, encoding='utf-8')))
    out = {}
    for fy in sorted({int(r['fy']) for r in rows}):
        rs = sorted((r for r in rows if int(r['fy']) == fy), key=lambda r: (r['eff_date'], int(r['source_row'])))
        soy = [r for r in rs if r['src'] == 'SOY']
        if len(soy) != 1:
            fail('FY%d: %d opening rows in the cashbook' % (fy, len(soy)))
        op = Decimal(soy[0]['amount'])
        mv = [Decimal(r['amount']) for r in rs if r['src'] != 'SOY']
        bal, low, low_d = ZERO, None, None
        for r in rs:
            bal += Decimal(r['amount'])
            if low is None or bal < low:
                low, low_d = bal, r['eff_date']
        out[fy] = dict(opening=op, receipts=sum((v for v in mv if v > 0), ZERO),
                       payments=sum((v for v in mv if v < 0), ZERO),
                       journals=sum((Decimal(r['amount']) for r in rs if r['src'] == 'GEN'), ZERO),
                       closing=bal, low=low, low_date=low_d)
    ys = sorted(out)
    for a, b in zip(ys, ys[1:]):
        if out[b]['opening'] != out[a]['closing']:
            fail('the cashbook closes FY%d at %s and opens FY%d at %s' % (a, out[a]['closing'], b, out[b]['opening']))
    return out


def year_ends(cash):
    """The printed year-end figures after FY2023, each from its own document, and the
    assertions that say what basis each is on."""
    r24 = [r for r in csv.DictReader(open(PR24, encoding='utf-8'))
           if r['fund_number'] == FUND and r['column'] == 'fund_balance']
    if len(r24) != 1:
        fail('the FY2024 annual report read holds %d fund 1301 balances' % len(r24))
    fb24 = Decimal(r24[0]['value'])
    if AR25_LINE not in open(AR25, encoding='utf-8').read():
        fail('the FY2025 annual report no longer prints %r' % AR25_LINE)
    ar25 = Decimal(AR25_LINE.split('$')[1].strip().replace(',', ''))
    if ar25 != cash[2025]['closing']:
        fail('the FY2025 annual report’s 1301 figure is no longer the cashbook’s close')
    import openpyxl
    ws = openpyxl.load_workbook(FUNDS26, data_only=True)['Athletics Revolving']
    if not (str(ws['A5'].value).startswith('Beginning Undesignated Fund Balance (7/1/25)')
            and str(ws['A8'].value).startswith('Ending Fund Balance / Cash (6/30/26)')):
        fail('school-funds-fy26.xlsx no longer labels C5 and C8 as before')
    fb25 = Decimal(str(ws['C5'].value)).quantize(Decimal('0.01'))
    fb26 = Decimal(str(ws['C8'].value)).quantize(Decimal('0.01'))
    payable = ar25 - fb25
    if '$%s' % format(payable, ',.2f') not in str(ws['A26'].value):
        fail('the district file no longer names the cash-to-fund-balance gap in A26')
    return dict(fb24=fb24, ar25=ar25, fb25=fb25, fb26=fb26, payable=payable)


def said_file(rel, quote, board):
    """A quotation asserted verbatim against our copy of the minutes."""
    text = ' '.join(open(os.path.join(MINUTES, rel), encoding='utf-8').read().split())
    if ' '.join(quote.split()) not in text:
        fail('quote is not verbatim in %s: %r' % (rel, quote))
    date = re.search(r'(\d{4}-\d{2}-\d{2})', rel).group(1)
    return dict(quote=quote, board=board, date=date, href='/docs/minutes/text/' + rel,
                label='%s minutes, %s' % (board, date))


def sources():
    return [
        dict(label='Town of Lunenburg — annual town reports FY2011–FY2023, special revenue '
                   'funds schedule (Chapter 658 School Athletics), read and tied in '
                   'special-revenue-read.csv', href='/docs/data/special-revenue-read.csv'),
        dict(label='Lunenburg Public Schools — Proposed FY19 MSHS Athletic Budget (both sides, '
                   'FY14–FY19)', href='/docs/district-budget/docs/fy19-proposed-athletics-budget.pdf'),
        dict(label='Lunenburg Budget Project — munis-school-journal.csv (every school journal '
                   'line, FY2023–FY2026, our redacted reading of the Town’s Report 1)',
             href=J.PUBLIC['journal']),
        dict(label='Town of Lunenburg — school special funds, period 13, FY2026 (fund 1301 tied '
                   'to the cent every year)',
             href='/docs/town-ledgers/expenses/glytdbud-expense-fy2026-p13-special-school.xlsx'),
        dict(label='Town of Lunenburg — how the journal reached us, and what was redacted',
             href='/docs/town-ledgers/account-details/PROVENANCE-fy2023-fy2026-p13-school-journal.md'),
        dict(label='Town of Lunenburg — FY2024 annual town report, Special Revenue Fund Balance Detail, '
                   'page 23 (fund 1301 at 30 June 2024)',
             href='/docs/town-annual-reports/docs/4132-fy-2024-annual-town-report.pdf'),
        dict(label='Town of Lunenburg — FY2025 annual town report, Special Revenue Fund Balance Detail '
                   '(fund 1301 at 30 June 2025, unaudited)',
             href='/docs/town-annual-reports/docs/4130-fy-2025-annual-town-report.pdf'),
        dict(label='Lunenburg Public Schools — school-funds-fy26.xlsx (the district’s FY2026 '
                   'reconciliation of the athletics revolving fund)',
             href='/docs/budget-workbooks/school-funds-fy26.xlsx'),
        dict(label='Lunenburg Budget Project — fund-1301-cash-journal.csv (our redacted extract of the '
                   'Town’s MUNIS journal exports for fund 1301, every posting; the originals are '
                   'withheld: they name a family)', href='/docs/data/fund-1301-cash-journal.csv'),
    ]


def build():
    ar = annual()
    f19, sec_line = fy19()
    t, moves = journal()
    cash = cashbook()
    ye = year_ends(cash)
    U = lambda v: figure(float(v), C.usd(v))
    F = lambda v: figure(v, C.fy(v), 'fiscal year')
    f = LD.fact
    con = TP.conclusion
    fv = float
    a0, a1 = min(ar), max(ar)
    j0, j1 = Y[0], Y[-1]
    if a1 != j0:
        fail('the annual reports should end where the journal starts (FY%d); they end FY%d' % (j0, a1))

    # FY2023 is in both: the annual report counts $X more on BOTH sides -- refunds gross, by
    # our reading -- and the net change is the same. Asserted, never assumed.
    d_in = ar[j0]['receipts'] - t[('in', j0)]
    d_out = ar[j0]['disb'] - t[('town_out', j0)]
    if d_in != d_out or d_in < 0:
        fail('FY%d: the annual report and the journal no longer differ by the same amount on '
             'both sides (%s, %s)' % (j0, d_in, d_out))

    mv = {y: sum((m['amount'] for m in moves if m['fy'] == y), ZERO) for y in Y}
    if mv[2025] <= t[('in', 2025)]:
        fail('the finding says FY2025’s moves exceeded its fees')
    in25, mv25, mv24 = t[('in', 2025)], mv[2025], mv[2024]
    peak_y = max(ar, key=lambda y: ar[y]['carried'])
    low_y = min((y for y in ar if y >= 2019), key=lambda y: ar[y]['disb'])
    sec = {y: f19[y]['Athletic Secretary'][0] for y in (2015, 2016, 2017)}
    pay26 = t[('pay', j1)]
    in26, out26 = t[('in', j1)], t[('out', j1)]
    if any(t[('payee', j1, p)] for p in PAYEES['bus'][1]):
        fail('the fund paid a bus company in FY%d; conclusion 4 says it paid none' % j1)
    if not in26 > out26:
        fail('the fund no longer takes in more than it spends in FY%d' % j1)

    px = {}
    px['in25'] = f('in25', U(in25), 'fees paid into the fund, FY2025', 'published', {'y': F(2025)})
    px['mv25'] = f('mv25', U(mv25), 'moved onto the town budget, FY2025', 'published', {'y': F(2025)})
    px['mv24'] = f('mv24', U(mv24), 'moved onto the town budget, FY2024', 'published', {'y': F(2024)})
    px['in26'] = f('in26', U(in26), 'fees paid into the fund, FY2026', 'published', {'y': F(j1)})
    px['out26'] = f('out26', U(out26), 'spent from it, FY2026', 'published', {'y': F(j1)})
    px['pay26'] = f('pay26', U(pay26), 'in pay, FY2026', 'published', {'y': F(j1)})
    px['peak'] = f('peak', U(ar[peak_y]['carried']), 'held at the end of %s, its most' % C.fy(peak_y),
                   'published', {'y': F(peak_y)})
    px['low'] = f('low', U(ar[low_y]['disb']), 'spent in %s, the lowest year' % C.fy(low_y),
                  'published', {'y': F(low_y)})
    px['sec'] = f('sec', U(sec[2017]), 'to the athletic secretary, FY2017 (district)', 'published',
                  {'y': F(2017)})
    px['how'] = f('how', None, 'what the fund bought, FY2020 to FY2022', 'gap', closes=HOW_CLOSES,
                  allow=D('FY2020 FY2022', HOW_CLOSES))
    px['who'] = f('who', None, 'whose pay the fund’s payroll lines were', 'gap', closes=WHO_CLOSES,
                  allow=D(WHO_CLOSES))
    px['memo'] = f('memo', None, 'what the memos behind the moves say', 'gap', closes=MEMO_CLOSES,
                   allow=D(MEMO_CLOSES))

    concl = [
        con('In FY2025 journal entries moved %s of the fund’s costs onto the town budget, against '
            '%s of fees.' % (C.usd(mv25), C.usd(in25)), ['mv25', 'in25', 'mv24', 'memo'],
            {'a': U(mv25), 'b': U(in25), 'y': F(2025)},
            so='The town budget’s athletics lines that year carry costs first charged to the fees.',
            notes=[step('FY2024 had one such entry, %s, onto regular school transportation. Every '
                        'move runs fee fund to town budget; none runs back.' % C.usd(mv24),
                        'measured', {'a': U(mv24), 'y': F(2024)}),
                   step('Each entry cites a memo the journal does not hold.', 'gap',
                        allow=D(MEMO_CLOSES), closes=MEMO_CLOSES)]),
        con('The fund held %s at the end of %s, then spent %s in %s.'
            % (C.usd(ar[peak_y]['carried']), C.fy(peak_y), C.usd(ar[j0]['disb']), C.fy(j0)),
            ['peak', 'low'], {'a': U(ar[peak_y]['carried']), 'y': F(peak_y),
                              'b': U(ar[j0]['disb']), 'z': F(j0)},
            so='The balance built up in FY2021 and FY2022 and was spent down before the FY2025 moves.',
            so_figures={'y': F(2025), 'a': F(2021), 'b': F(2022)},
            notes=[step('%s, the lowest year, the fund spent %s. That the pandemic caused it is '
                        'our reading; the report gives no reason.' % (C.fy(low_y), C.usd(ar[low_y]['disb'])),
                        'hypothesis', {'y': F(low_y), 'a': U(ar[low_y]['disb'])})]),
        con('The district’s FY19 document pays an athletic secretary from the fund, %s in FY2017, '
            'and no director.' % C.usd(sec[2017]), ['sec', 'pay26', 'who'],
            {'a': U(sec[2017]), 'y': F(2017)}, allow=('FY19',),
            so='The fees have paid athletics staff in every year itemised; the director is not one of them.',
            notes=[step('Its line reads “%s”.' % sec_line, 'measured',
                        allow=tuple(re.findall(r'\d[\d,]*', sec_line))),
                   step('From %s the journal records pay as one line per pay run, with no name: '
                        '%s in %s.' % (C.fy(j0), C.usd(pay26), C.fy(j1)), 'gap',
                        {'y': F(j0), 'a': U(pay26), 'z': F(j1)}, allow=D(WHO_CLOSES),
                        closes=WHO_CLOSES)]),
        con('In %s the fund took in %s and spent %s.' % (C.fy(j1), C.usd(in26), C.usd(out26)),
            ['in26', 'out26', 'how'], {'y': F(j1), 'a': U(in26), 'b': U(out26)},
            so='That year the fund paid no bus company, and its fees exceeded its spending.',
            notes=[step('Every year from %s ties to the cent to the Town’s period-13 report. %s '
                        'differs from the annual report by %s on both sides, so the net agrees.'
                        % (C.fy(j0), C.fy(j0), usd2(d_in)), 'measured',
                        {'y': F(j0), 'a': figure(float(d_in), usd2(d_in))}, allow=('13',))]),
    ]

    # ---- the signature: money in and money out, every year held
    years = list(range(a0, j1 + 1))
    xs = [C.fy(y) for y in years]
    rin = [fv(ar[y]['receipts']) if y in ar and y < j0 else fv(t[('in', y)]) for y in years]
    rout = [fv(ar[y]['disb']) if y in ar and y < j0 else fv(t[('out', y)]) for y in years]
    before = [fv(t[('booked', y)]) if y in Y and mv[y] else None for y in years]
    chart = series('Money in and money out', xs, [
        TP.track('Fees and other receipts', rin, 'published'),
        TP.track('Spent', rout, 'published'),
        TP.track('Charged before the moves', before, 'published')], '$', kind='line',
        note='%s–%s: annual town reports. %s on: the school journal, spent after the moves.'
             % (C.fy(a0), C.fy(j0 - 1), C.fy(j0)))

    # ---- what it bought: the FY19 document's items, then the journal's groups
    f19y = [y for y in sorted(f19) if y <= 2017]
    items = [('Athletic Coaches', 'Coaches'), ('Athletic Transportation', 'Buses'),
             ('Athletic Expenses/Supplies', 'Supplies'), ('Athletic Secretary', 'Athletic secretary')]
    also = [series('What the fees bought, district', [C.fy(y) for y in f19y],
                      [TP.track(n, [fv(f19[y].get(k, (ZERO,))[0]) for y in f19y], 'published')
                       for k, n in items], '$', kind='bar', stacked=True,
                      note='Actuals, the fee-fund side of the district’s FY19 athletics document. '
                           'It prints no director line.')]
    bands = [('Pay', ['pay']), ('Buses', ['bus']), ('ArbiterSports', ['arb']),
             ('Uniforms, video and gear', ['gear', 'video']), ('Everything else', ['site', 'org', 'rest'])]
    if sorted(g for _, gs in bands for g in gs) != sorted(ORDER):
        fail('the bands do not hold every group once')
    also.append(series('What the fees were charged with', [C.fy(y) for y in Y],
                          [TP.track(n, [fv(sum((t[('g', y, g)] for g in gs), ZERO)) for y in Y], 'published')
                           for n, gs in bands], '$', kind='bar', stacked=True,
                          note='As booked, before any move onto the town budget. Payees grouped by '
                               'what each organisation is, our grouping.'))

    # ---- table 1: the moves, both ends
    rows = []
    for m in moves:
        for desc, amt in m['to']:
            rows.append(dict(label='%s · %s' % (long_date(m['date']), m['memo'].title()),
                             cells=[m['journal'] + (' (reversal)' if m['source'] == 'GRV' else ''),
                                    ', '.join(x.title() for x in m['frm']), desc.title(), usd2(amt)]))
    rows.append(dict(label='Net, every entry', cells=['', '', '', usd2(sum((m['amount'] for m in moves), ZERO))]))
    if len(rows) > TP.TABLE_ROWS_MAX:
        fail('the moves table runs past %d rows' % TP.TABLE_ROWS_MAX)
    t_moves = dict(caption='Every move between the fee fund and the town budget', views=[dict(
        key='moves', label='Both ends', kinds=['stated', 'stated', 'stated', 'ledger'],
        columns=['Journal entry', 'Taken off (fee fund)', 'Put on (town budget)', 'Amount'], rows=rows,
        note='Read off the school journal. The FY2024 entry was made, reversed and made again; all '
             'three are shown.')])

    # ---- table 2: the fund, year by year
    cols = ['Money in', 'Spent', 'Moved to town budget', 'Held at year end', 'Source']
    yr = []
    for y in years:
        if y in ar and y < j0:
            yr.append(dict(label=C.fy(y), cells=[usd2(ar[y]['receipts']), usd2(ar[y]['disb']), '—',
                                                usd2(ar[y]['carried']), 'Annual report p%s' % ar[y]['page']]))
        else:
            held, where = {
                j0: (ar[j0]['carried'], '; year end, annual report p%s' % ar[j0]['page']),
                2024: (ye['fb24'], '; year end, FY2024 annual report p23'),
                2025: (ye['ar25'], '; year end, FY2025 annual report: the fund’s cash'),
                2026: (ye['fb26'], '; year end, the district’s FY2026 reconciliation'),
            }.get(y, (None, ''))
            yr.append(dict(label=C.fy(y), cells=[usd2(t[('in', y)]), usd2(t[('out', y)]),
                                                usd2(-t[('moved', y)]) if t[('moved', y)] else '—',
                                                usd2(held) if held is not None else '—',
                                                'School journal' + where]))
    cy = sorted(cash)
    crow = [dict(label=C.fy(y), cells=[usd2(cash[y][k]) for k in ('opening', 'receipts', 'payments',
                                                                 'journals', 'closing', 'low')])
            for y in sorted(cy, reverse=True)]
    t_years = dict(caption='The fund, every year held', views=[dict(
        key='years', label='By year', kinds=['ledger', 'ledger', 'ledger', 'ledger', 'stated'],
        columns=cols, rows=yr,
        note='Year ends on three bases: fund balance (annual reports to FY2024), cash (the FY2025 '
             'report, %s above the fund balance the district opens FY2026 with: unpaid bills), and '
             'the district’s own FY2026 sheet.' % usd2(ye['payable']))])
    t_cash = dict(caption='The fund’s cash, by year', views=[
        dict(key='cash', label='Cash, the fund’s cashbook', total_col=4, kinds=['ledger'] * 6,
             columns=['Opening cash', 'Money in', 'Paid out', 'Of money in: journal entries',
                      'Closing cash', 'Lowest cash in the year (ours)'], rows=crow,
             note='Our redacted reading of the Town’s MUNIS journal exports for fund 1301. Each year’s '
                  'close is the next year’s printed opening. The latest year stops on %s. The lowest '
                  'cash is our running total, in effective-date order.' % long_date(max(
                      r['eff_date'] for r in csv.DictReader(open(CASH, encoding='utf-8')))))])

    # ---- table 3: who the fund paid, journal years, organisations only
    pcols = [C.fy(y) for y in Y] + ['Four years']
    four = lambda k: sum((t[(k[0], y) + k[1:]] for y in Y), ZERO)
    payees = {k[2] for k in t if k[0] == 'payee'}
    prow, shown = [], []
    for g in ORDER:
        if g in ('pay', 'rest'):
            continue
        for p in sorted((p for p in payees if GROUP_OF[p] == g), key=lambda p: (-four(('payee', p)), p)):
            if four(('payee', p)) < NAMED_MIN:
                continue
            shown.append(p)
            prow.append(dict(label='%s · %s' % (p, LABEL[g]),
                             cells=[C.usd(t[('payee', y, p)]) for y in Y] + [C.usd(four(('payee', p)))]))
    rest = {y: t[('booked', y)] - t[('pay', y)] - sum((t[('payee', y, p)] for p in shown), ZERO) for y in Y}
    prow.append(dict(label='Every other payee, people among them',
                     cells=[C.usd(rest[y]) for y in Y] + [C.usd(sum(rest.values(), ZERO))]))
    prow.append(dict(label='Pay', cells=[C.usd(t[('pay', y)]) for y in Y] + [C.usd(four(('pay',)))]))
    prow.append(dict(label='Charged to the fees', cells=[C.usd(t[('booked', y)]) for y in Y]
                     + [C.usd(four(('booked',)))]))
    prow.append(dict(label='Moved onto the town budget', cells=[C.usd(t[('moved', y)]) for y in Y]
                     + [C.usd(four(('moved',)))]))
    prow.append(dict(label='Spent', cells=[C.usd(t[('out', y)]) for y in Y] + [C.usd(four(('out',)))]))
    t_payees = dict(caption='Who the fees paid, by year', views=[dict(
        key='payees', label='By payee', total_col=4, kinds=['ledger'] * len(pcols), columns=pcols,
        rows=prow, note='Organisations paid %s or more over four years, as the journal prints them. '
                        'People are never named.' % C.usd(NAMED_MIN))])

    page = TP.page(
        TID,
        'In FY2025 the town budget took %s of the fee fund’s costs by journal entry, against %s of fees.'
        % (C.usd(mv25), C.usd(in25)),
        list(px.values()), concl, chart=chart, also=also, tables=[t_moves, t_years, t_cash, t_payees],
        said=[said_file('school-committee/2025-09-03-minutes-7385.txt',
                        'in the past 6 months or so there have been things that have come up, '
                        'athletic revolving fund was running in a deficit, I was told over '
                        '$100,000', 'School Committee'),
              said_file('finance-committee/2025-07-08-minutes-7295.txt',
                        'She emphasized the need for all revolving accounts to have a budget, '
                        'even if they are not town meeting revolving accounts, to prevent '
                        'negative balances from going unnoticed.', 'Finance Committee')],
        sources=sources(), finding_figures={'a': U(mv25), 'b': U(in25), 'y': F(2025)})
    page['not_established'] = [
        'What the fund paid for in FY2020 to FY2022. The annual reports print one total a year.',
        'Whose pay the fund’s payroll lines were. The journal records one line per account per pay run.',
        'Why each move was made. The entries cite memos the journal does not hold.',
        'One basis for every year end: the FY2025 report prints the fund’s cash, the others a fund balance.',
    ]
    if [TID] != MINE:
        fail('the registry gives this script %s; it builds %s' % (MINE, [TID]))
    return {TID: page}


def outputs():
    return {os.path.join(OUT_DIR, '%s.json' % t):
            json.dumps(p, indent=1, ensure_ascii=False, sort_keys=True) + '\n'
            for t, p in build().items()}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--check', action='store_true')
    a = ap.parse_args()
    outs = outputs()
    if a.check:
        stale = [os.path.relpath(p, ROOT) for p, text in outs.items()
                 if not os.path.exists(p) or open(p, encoding='utf-8').read() != text]
        if stale:
            print('STALE %s' % ', '.join(stale), file=sys.stderr)
            return 1
        print('fee fund history: %d output current' % len(outs))
        return 0
    for p, text in outs.items():
        os.makedirs(os.path.dirname(p), exist_ok=True)
        tmp = p + '.tmp'
        with open(tmp, 'w', encoding='utf-8', newline='\n') as fh:
            fh.write(text)
        os.replace(tmp, p)
        print('wrote %s' % os.path.relpath(p, ROOT))
    return 0


if __name__ == '__main__':
    sys.exit(main())
