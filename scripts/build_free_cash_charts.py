#!/usr/bin/env python3
"""The free-cash drill-in's series, pre-rendered from the DLS proof.

    python3 scripts/build_free_cash_charts.py            # write it
    python3 scripts/build_free_cash_charts.py --check    # fail if it is stale

WHAT THIS IS. `free_cash_proof` is the Division of Local Services' own **proof** of free
cash for Lunenburg and eight comparable towns, 2021-2025 -- 630 rows out of nine
workbooks. A proof is not a summary: it prints the components and it prints the total, so
the arithmetic by which free cash comes into existence can be SHOWN to a resident rather
than asserted. That is the whole reason this file exists.

THE CROSS-FOOT, AND WHAT IT ACTUALLY TIES TO. The eleven component rows sum, to the cent,
in all 45 town-years -- but they sum to `Identified Free Cash July 1,` (row 17) and NOT to
`Current Year Calculation` (row 5), which is the figure DLS certifies and the Town quotes.
Identified and certified differ in all 45 town-years, by $361,912 for Lunenburg in 2025,
and the workbook prints no line reconciling them. That is a finding, it is carried in
`crossfoot`, and it is stated on the page rather than smoothed. Rule 7 applies to it: we
hold no DLS document giving the reason and do not guess at one.

RULE 13, TWICE.

  1. `free_cash_proof.source_ref` is `Sheet1!A<row>` for every row of every year -- the
     LABEL cell in column A, not the cell the amount came from. The amount for 2025 sits
     in column F. So this generator derives the VALUE coordinate itself (`Sheet1!F10`) and,
     when the workbook is on disk, OPENS IT and asserts that the cited cell holds the cited
     amount. A coordinate quoted that is not the coordinate the number came from is exactly
     the defect rule 13 is about, and the fix is to cite the real one.
  2. Every figure `model/freecash.py` states as a constant -- CERTIFIED, IDENTIFIED,
     UNSPENT_2025, UNSPENT_AVG_2021_24 and the nine typed PEER_MULTIPLES -- is recomputed
     here from the table and asserted. They were typed; now they are checked.

RULE 11 IS THE SUBJECT AGAIN, POINTED AT A RESERVE. A free cash figure is net of
everything and one-time by construction, and the proof shows why: nearly half of five
years of it is money appropriated and not spent, and another 11% is last year's free cash
that was never appropriated. Free cash is not a savings account; it is the residue of
estimating.

THE TOWN'S OWN PROOF, BESIDE THE STATE'S. The annual town report prints an `Undesignated
Fund Balance Roll-forward` -- the same add/deduct worksheet, produced by the Town rather
than by DLS, for the same 30 June date. It is read here from two DIFFERENT instruments and
both are named on the page, because rule 13 says the instrument is part of the finding:
FY2024 comes off the OCR of the rendered page, FY2025 off the PDF's own text layer. Each
year's steps are cross-footed against the two totals the page itself prints, and FY2024's
closing balance must equal FY2025's opening one.

It does NOT reconcile to the DLS proof, and nothing published says why. Three figures
describe the same balance sheet on the same date -- the Town's undesignated fund balance,
DLS's identified free cash, and DLS's certified free cash -- and no document in this
archive prints a line between any pair of them. That gap is carried in `versus_dls` and
registered, not smoothed.

WHAT IT REFUSES TO WRITE ON. Nine checks, each guarding a join that could match nothing:
  1. coverage must be rectangular -- every town, every year, fourteen rows;
  2. components must sum to `Identified Free Cash July 1,` in every town-year;
  3. `Current Year Calculation` in year N must equal `Free Cash Certified Prior Year` in
     year N+1, for every town -- the source's own second identity;
  4. the workbook, where present, must hold every cited amount at the cited coordinate;
  5. every quote from the meeting archive must be found verbatim in the file it names;
  6. `money_gaps` must return the rows this page quotes;
  7. each roll-forward's Add/Deduct steps must sum to the subtotal the page prints;
  8. ...and on to the `Current Year Undesignated Fund Balance` the page prints;
  9. FY2024's closing balance must be FY2025's opening balance.
"""
import argparse
import csv
import json
import re
import os
import sqlite3
import statistics
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, 'model'))
DB = os.path.join(ROOT, 'sources/data/lunenburg.db')
OUT = os.path.join(ROOT, 'fy28/public/data/free-cash.json')
WORKBOOKS = os.path.join(ROOT, 'sources/state-dls')
MINUTES = os.path.join(ROOT, 'sources', 'meetings', 'text')

TOWN = 'Lunenburg'
CENT = 0.005

CERT = 'Current Year Calculation'
PRIOR = 'Free Cash Certified Prior Year'
IDENT = 'Identified Free Cash July 1,'
UNSPENT = 'Add Unencumbered/Unexpended Appropriations (CL#11)'
RECEIPTS = 'Excess/Shortfall Local Receipts (CL#6)'
AID = 'Excess/Shortfall Cherry Sheet Receipts (CL#8)'
RECYCLED = 'Add Prior Year Free Cash Not Appropriated (CL#12)'

# The four components that carry the story, and what each one MEANS in a resident's
# words. Everything else is a residual and is drawn achromatic for that reason -- the
# same treatment `build_athletics_charts.py` gives its 'Everything else'.
# THREE NAMED PARTS AND A RESIDUAL, and the residual is defined ONCE. `other` is the
# identified total minus exactly these three, which means the stack sums to a figure DLS
# prints and the same definition holds on the peer chart. `Add Prior Year Free Cash Not
# Appropriated (CL#12)` is INSIDE that residual and is reported separately for the table
# beside it -- it is last year's free cash arriving again, which is worth a reader's
# attention and is not a fourth source of money.
NAMED = [
    (UNSPENT, 'underspend', 'Money appropriated and not spent'),
    (RECEIPTS, 'receipts', 'Local receipts above the estimate'),
    (AID, 'aid', 'State aid against the estimate'),
]
PART_ORDER = ['underspend', 'receipts', 'aid', 'other']

# What the town said, in its own meetings, about the two things this page measures.
# Rule 15a: the archive is searched for what people said about a category in the same
# years the figures cover, and the quote goes BESIDE the measurement. Each is verified
# verbatim against the file named -- `minutes_decisions.py`'s pattern, and the reason is
# the same: a quote typed from memory is a figure typed into prose wearing a costume.
BASE = 'https://www.lunenburgma.gov/AgendaCenter/ViewFile/Minutes/'
SITE = 'https://lunenburgbudgetproject.org/docs/minutes/text/'
SAID = [
    dict(key='conservative',
         board='Finance Committee', date='2026-05-28',
         path='finance-committee/2026-05-28-minutes-7828.txt', suffix='_05282026-7828',
         agenda='b. Local Receipts Trends',
         quote='Ana Lockwood states for the last 10 years the town has not collected less '
               'than 4 million dollars in local receipts. This year’s estimate is 3.535 '
               'million. Ana Lockwood suggests not being so conservative based on the data.',
         beside='receipts'),
    dict(key='review_policy',
         board='Finance Committee', date='2026-05-28',
         path='finance-committee/2026-05-28-minutes-7828.txt', suffix='_05282026-7828',
         agenda='c. Discuss FY27 Budget process and improvement ideas',
         quote='Chris Menard suggests having an analyst review all of the town’s policies '
               'on free cash and estimates.',
         beside='receipts'),
    dict(key='budgeted_conservatively',
         board='Select Board', date='2025-03-11',
         path='select-board/2025-03-11-minutes-7096.txt', suffix='_03112025-7096',
         agenda='FY26 Preliminary Budget',
         quote='Local receipts are budgeted conserva',
         beside='receipts'),
    dict(key='missing_out',
         board='School Committee', date='2025-09-03',
         path='school-committee/2025-09-03-minutes-7385.txt', suffix='_09032025-7385',
         agenda='d. FY 25 Budget',
         quote='Ms. Squier shared some of the things her students are missing out on and '
               'how hard it has made it, which is disappointing and frustrating based on '
               'an unknown surplus of money that was not spent and instead was given back '
               'to the town. Mr. Yvon shares some of the positions that were cut and left '
               'unfilled and the ability to meet all the students ’ needs is harder. Mr. '
               'Santry shares we had drastic reductions in services to students. We really '
               'could have used those funds to keep some of the positions.',
         beside='underspend'),
    dict(key='not_lost',
         board='School Committee', date='2025-09-03',
         path='school-committee/2025-09-03-minutes-7385.txt', suffix='_09032025-7385',
         agenda='',
         quote='With respect to the unspent money, the money is not lost. The money is '
               'still there, it is free cash and it is going to roll over. It is going to '
               'be appropriated in one fashion or another.',
         beside='underspend'),
    dict(key='stopgap',
         board='Select Board', date='2025-03-11',
         path='select-board/2025-03-11-minutes-7096.txt', suffix='_03112025-7096',
         agenda='FY26 budget gap',
         quote='The options discussed included a multi-year override or utilizing one-time '
               'funds specifically free cash as a stopgap, however these funds should not '
               'be used in an operating budget, but as a one-time appropriation.',
         beside='policy'),
    dict(key='certified',
         board='Finance Committee', date='2026-03-05',
         path='finance-committee/2026-03-05-minutes-7682.txt', suffix='_03052026-7682',
         agenda='Announcements',
         quote='free cash has been certi',
         beside='certified'),
]

# The gap-register rows this page rests on, quoted BY KEY. If one is renamed out from
# under it the page would render an empty box, so the read is asserted.
GAP_KEYS = [
    ('money_in', 'Why DLS certifies a different amount from the free cash its own proof identifies'),
    ('money_in', 'Which departments’ unspent appropriations produced the free cash'),
    ('money_in', 'How much of free cash is the schools’'),
    ('money_in', 'Whether Lunenburg’s free cash is unusual for a town its size'),
    ('money_in', 'Why the Town’s undesignated fund balance and DLS’s free cash differ'),
    # The PEOPLE side of the same question, and it was already in the register before this
    # page existed. A large unexpended-appropriations line can be prudence, a project that
    # slipped, a line over-budgeted from the start, or a post nobody filled -- and the
    # fourth of those is a question about people rather than about money. Quoted here
    # rather than restated, because rule 7c says the registry outranks the page.
    ('people', 'Whether a budgeted position was filled'),
]

# Rule 2 reaches the register too. These rows state figures in prose, so the figures are
# recomputed here and the row is required to still contain them -- a gap whose numbers
# have drifted from the data is a grievance rather than a records request.
GAP_FIGURES = {
    'Why DLS certifies a different amount from the free cash its own proof identifies':
        ['identified_minus_certified_2025'],
    'Why the Town’s undesignated fund balance and DLS’s free cash differ':
        ['undesignated_2025', 'identified_2025', 'certified_2025'],
}


# --------------------------------------------------------------------------------------
# The TOWN's own version of the same arithmetic, out of the annual town report.
#
# `Undesignated Fund Balance Roll-forward` is an add/deduct worksheet ending in a figure
# the page itself calls `Current Year Undesignated Fund Balance`, with a printed PROOF
# line under it. It is the same 30 June balance sheet DLS certifies free cash from, stated
# by the Town instead of by the state.
#
# TWO DIFFERENT INSTRUMENTS, and rule 13 says which one read a figure is part of the
# figure. FY2025's page carries a text layer and is read from it. FY2024's does not, and
# is read from the OCR of the rendered page -- token positions, paired label to amount by
# row. Both are named on the page.
#
# The labels below are QUOTED from the documents, not paraphrased, and a label that stops
# appearing fails the build. The amounts are never typed: they are read, and then checked
# against the two totals the page prints itself.
MONEY = re.compile(r'\d{1,3}(?:,\d{3})*\.\d{2}')
MONEY_ONLY = re.compile(r'^\d{1,3}(?:,\d{3})*\.\d{2}$')

ROLLFORWARD = [
    dict(fy=2024, document='4132-fy-2024-annual-town-report.pdf',
         reader='ocr',
         path='sources/town-budget/ocr/4132-fy-2024-annual-town-report.tsv',
         page='30', printed_page='30', pages_note='PDF page 30; the page prints 30 in its '
                                                  'own footer',
         as_of='June 30, 2024',
         url='/docs/town-annual-reports/docs/4132-fy-2024-annual-town-report.pdf',
         opening='Beginning Undesignated Fund Balance',
         subtotal='Prior Year Total Fund Balance',
         total='Current Year Undesignated Fund Balance',
         above=[('Prior Year Reserved for Encumbrances', 1),
                ('Prior Year Reserved for / Reserved for Debt Service', 1),
                ('Prior Year Reserved for / Premiums Reserved for Debt', 1),
                ('Prior Year Reserved for Expenditures', 1)],
         below=[('Current Year Reserved for Encumbrances', -1),
                ('Current Year Reserved for Expenditures', -1),
                ('Current Year Reserved for / Premium Reserved for Debt', -1),
                ('Add: / Current Year Revenue Closeouts', 1),
                ('Current Year Expenditure Closeouts', -1)]),
    dict(fy=2025, document='4130-fy-2025-annual-town-report.pdf',
         reader='text',
         path='sources/town-annual-reports/text/4130-fy-2025-annual-town-report.txt',
         page='31-33', printed_page='27-29',
         pages_note='PDF pages 31–33; the pages print 27–29 in their own footers, which is '
                    'what the report’s table of contents names',
         as_of='June 30, 2025',
         url='/docs/town-annual-reports/docs/4130-fy-2025-annual-town-report.pdf',
         opening='Beginning undesignated Fund Balance',
         subtotal='Prior Year Total Fund Balance',
         total='Current Year Undesignated Fund Balance',
         above=[('Prior Year Reserved for Encumbrance', 1)],
         below=[('Current Year Reserved for Encumbrance', -1),
                ('Current Reserved for Expenditures', -1),
                ('Current Year Reserved for ATM ART #7 5/4/2025 ATM', -1),
                ('Current Year Revenue Closeouts', 1),
                ('Current Year Expenditure Closeouts', -1)]),
]


def _read_ocr_page(path, page):
    """Every amount on one OCR'd page, paired to the label tokens on its own row."""
    got = {}
    with open(os.path.join(ROOT, path), encoding='utf-8') as fh:
        rd = csv.reader(fh, delimiter='\t')
        next(rd)
        toks = [dict(x=float(r[1]), y=float(r[2]), text=r[6]) for r in rd if r[0] == page]
    if not toks:
        fail(f'{path} has no tokens on page {page} — the page this reads was not OCR’d, '
             'and an empty page reads exactly like a page with nothing on it')
    amounts = [t for t in toks if MONEY_ONLY.match(t['text'])]
    labels = [t for t in toks if t['x'] < 0.5 and not MONEY_ONLY.match(t['text'])]
    for a in amounts:
        row = sorted((l for l in labels if abs(l['y'] - a['y']) < 0.008),
                     key=lambda l: l['x'])
        got[' / '.join(l['text'] for l in row)] = float(a['text'].replace(',', ''))
    return got


def _read_text_block(path, opening, total):
    """The roll-forward block out of the PDF's own text layer, label -> first amount after."""
    t = open(os.path.join(ROOT, path), encoding='utf-8').read()
    s = t.find('UNDESIGNATED FUND BALANCE ROLL-FORWARD')
    if s < 0:
        fail(f'{path} no longer contains the roll-forward heading')
    e = t.find('PROOF UNDESIGNATED FUND BALANCE', s)
    if e < 0:
        fail(f'{path}: the roll-forward has no PROOF line, which is the page’s own end')
    return t[s:e + 200]


def _amount_after(block, label, path):
    i = block.find(label)
    if i < 0:
        fail(f'{path} no longer prints the line {label!r} — this page quotes the document’s '
             'own labels, so a renamed line fails here rather than rendering a blank row')
    m = MONEY.search(block, i + len(label))
    if m is None:
        fail(f'{path}: no amount follows {label!r}')
    return float(m.group(0).replace(',', ''))


def rollforward():
    """The Town's own proof, read, cross-footed, and chained across two annual reports."""
    out = []
    for spec in ROLLFORWARD:
        if spec['reader'] == 'ocr':
            table = _read_ocr_page(spec['path'], spec['page'])

            def amount(label, _t=table, _p=spec['path']):
                if label not in _t:
                    fail(f'{_p}: the roll-forward has no row {label!r}')
                return _t[label]
        else:
            block = _read_text_block(spec['path'], spec['opening'], spec['total'])

            def amount(label, _b=block, _p=spec['path']):
                return _amount_after(_b, label, _p)

        opening = amount(spec['opening'])
        above = [dict(label=l, sign=g, amount=round(amount(l), 2)) for l, g in spec['above']]
        subtotal = amount(spec['subtotal'])
        below = [dict(label=l, sign=g, amount=round(amount(l), 2)) for l, g in spec['below']]
        total = amount(spec['total'])

        # CHECK 7 -- the first identity the page states about itself.
        got = opening + sum(r['sign'] * r['amount'] for r in above)
        if abs(got - subtotal) > CENT:
            fail(f'FY{spec["fy"]} roll-forward: the opening balance and the prior-year '
                 f'reservations come to {got:,.2f} against the {spec["subtotal"]!r} of '
                 f'{subtotal:,.2f} the page prints — the extract does not tie to its source')
        # CHECK 8 -- and the second, which is the whole worksheet.
        got = subtotal + sum(r['sign'] * r['amount'] for r in below)
        if abs(got - total) > CENT:
            fail(f'FY{spec["fy"]} roll-forward: the steps come to {got:,.2f} against the '
                 f'{spec["total"]!r} of {total:,.2f} the page prints')

        out.append(dict(
            fy=spec['fy'], document=spec['document'], url=spec['url'],
            reader=('the OCR of the rendered page' if spec['reader'] == 'ocr'
                    else 'the PDF’s own text layer'),
            reader_path=spec['path'], page=spec['page'], printed_page=spec['printed_page'],
            pages_note=spec['pages_note'], as_of=spec['as_of'],
            opening=dict(label=spec['opening'], amount=round(opening, 2)),
            above=above,
            subtotal=dict(label=spec['subtotal'], amount=round(subtotal, 2)),
            below=below,
            total=dict(label=spec['total'], amount=round(total, 2))))

    # CHECK 9 -- one year's closing balance is the next year's opening one, across two
    # separately published documents. The town's own chain, not ours.
    for a, b in zip(out, out[1:]):
        if abs(a['total']['amount'] - b['opening']['amount']) > CENT:
            fail(f'FY{a["fy"]} closes at {a["total"]["amount"]:,.2f} and FY{b["fy"]} opens '
                 f'at {b["opening"]["amount"]:,.2f} — the two annual reports no longer '
                 'chain, so they may not be drawn as one series')
    return out


def fail(msg):
    sys.exit(f'build_free_cash_charts: {msg}')


def col_letter(i):
    """Workbook column for the i-th year. B is 2021, F is 2025."""
    return chr(ord('B') + i)


def build():
    if not os.path.exists(DB):
        fail(f'{os.path.relpath(DB, ROOT)} is not here — run scripts/build_db.py')
    c = sqlite3.connect(DB)
    c.row_factory = sqlite3.Row
    raw = [dict(r) for r in c.execute(
        'SELECT town, year, line, amount, role, source_file, source_ref '
        'FROM free_cash_proof')]
    if not raw:
        fail('free_cash_proof is empty — the table this page is built from matched nothing')
    for r in raw:
        r['year'] = int(r['year'])
        r['amount'] = float(r['amount'])

    towns = sorted({r['town'] for r in raw})
    years = sorted({r['year'] for r in raw})
    if TOWN not in towns:
        fail(f'{TOWN} is not in free_cash_proof')

    # ---------------------------------------------------------------- CHECK 1: rectangular
    lines = [r['line'] for r in raw if r['town'] == TOWN and r['year'] == years[0]]
    for t in towns:
        for y in years:
            got = [r['line'] for r in raw if r['town'] == t and r['year'] == y]
            if sorted(got) != sorted(lines):
                fail(f'{t} {y} has {len(got)} rows against {len(lines)} for {TOWN} '
                     f'{years[0]} — the peer set is not comparable row for row, and a '
                     'comparison across a different set of lines per town is not a '
                     'comparison')

    def at(town, year, line):
        for r in raw:
            if r['town'] == town and r['year'] == year and r['line'] == line:
                return r['amount']
        fail(f'no row for {town} {year} {line!r}')

    comps = [ln for ln in lines
             if next(r['role'] for r in raw if r['line'] == ln) == 'component']
    comp_order = [ln for ln in lines if ln in comps]

    # ------------------------------------------- CHECK 2: the source's own cross-foot
    # Eleven components against the total the sheet prints. This is DLS's arithmetic,
    # not ours, and it is the reason the page may show a resident the derivation.
    crossfoot = []
    tied = 0
    for t in towns:
        for y in years:
            s = sum(at(t, y, ln) for ln in comps)
            ident = at(t, y, IDENT)
            cert = at(t, y, CERT)
            if abs(s - ident) > CENT:
                fail(f'{t} {y}: the components sum to {s:,.2f} against the workbook’s own '
                     f'{IDENT!r} of {ident:,.2f} — DLS’s proof no longer proves, so this '
                     'page may not claim that it does')
            tied += 1
            crossfoot.append(dict(town=t, year=y, components=round(s, 2),
                                  identified=round(ident, 2), certified=round(cert, 2),
                                  gap=round(ident - cert, 2),
                                  gap_share=round((ident - cert) / ident, 4) if ident else None))
    differ = [r for r in crossfoot if abs(r['gap']) > CENT]

    # --------------------------------- CHECK 3: the chain the source states about itself
    # `Current Year Calculation` in year N is `Free Cash Certified Prior Year` in N+1.
    chained = 0
    for t in towns:
        for a, b in zip(years, years[1:]):
            if abs(at(t, a, CERT) - at(t, b, PRIOR)) > CENT:
                fail(f'{t}: {CERT!r} for {a} is {at(t, a, CERT):,.2f} and {PRIOR!r} for '
                     f'{b} is {at(t, b, PRIOR):,.2f} — the proof no longer chains year to '
                     'year, so neither series may be drawn as one line')
            chained += 1

    # ------------------------------------- CHECK 4: the coordinate, against the workbook
    # `source_ref` is the label cell in column A for every year. The amount lives in the
    # year's own column, and that is what the page cites. Verified against the file when
    # the file is on disk (a fresh clone has the manifest and not the bytes).
    def ref(line, year):
        src = next(r['source_ref'] for r in raw
                   if r['town'] == TOWN and r['year'] == year and r['line'] == line)
        row = src.split('!A')[1]
        return f'Sheet1!{col_letter(years.index(year))}{row}', f'Sheet1!A{row}'

    checked_cells = 0
    wb_path = os.path.join(WORKBOOKS, f'free-cash-proof-{TOWN.lower()}.xlsx')
    if os.path.exists(wb_path):
        try:
            import openpyxl
        except ImportError:
            openpyxl = None
        if openpyxl is not None:
            ws = openpyxl.load_workbook(wb_path, data_only=True)['Sheet1']
            if (ws['A1'].value or '').strip() != TOWN:
                fail(f'{os.path.relpath(wb_path, ROOT)} A1 is {ws["A1"].value!r}, not '
                     f'{TOWN!r} — this is the mis-export shape PROVENANCE.md records')
            for ln in lines:
                for y in years:
                    coord, label_coord = ref(ln, y)
                    got = ws[coord.split('!')[1]].value
                    if got is None or abs(float(got) - at(TOWN, y, ln)) > CENT:
                        fail(f'{coord} holds {got!r}; the table says '
                             f'{at(TOWN, y, ln):,.2f} for {ln!r} in {y}')
                    lab = (ws[label_coord.split('!')[1]].value or '').strip()
                    if lab != ln:
                        fail(f'{label_coord} reads {lab!r}, not {ln!r}')
                    checked_cells += 1

    # ------------------------------------------------------- Lunenburg, the whole proof
    proof = []
    for ln in lines:
        role = next(r['role'] for r in raw if r['line'] == ln)
        proof.append(dict(
            line=ln, role=role,
            by_year=[dict(year=y, amount=round(at(TOWN, y, ln), 2), ref=ref(ln, y)[0])
                     for y in years]))

    # --------------------------------------------------- composition, Lunenburg by year
    composition = []
    for y in years:
        ident = at(TOWN, y, IDENT)
        parts = {key: round(at(TOWN, y, ln), 2) for ln, key, _ in NAMED}
        parts['other'] = round(ident - sum(parts.values()), 2)
        composition.append(dict(
            year=y, identified=round(ident, 2), certified=round(at(TOWN, y, CERT), 2),
            gap=round(ident - at(TOWN, y, CERT), 2),
            recycled=round(at(TOWN, y, RECYCLED), 2),
            **parts,
            shares={k: (round(v / ident, 4) if ident else None)
                    for k, v in parts.items()}))

    # Five years added together, which is the honest way to state what free cash IS --
    # and it is NOT five years of new money, because CL#12 is last year's free cash
    # arriving again. Said on the page beside the figure.
    total_ident = sum(r['identified'] for r in composition)
    by_line = []
    for ln in comp_order:
        s = sum(at(TOWN, y, ln) for y in years)
        by_line.append(dict(line=ln, amount=round(s, 2),
                            share=round(s / total_ident, 4) if total_ident else None))
    by_line.sort(key=lambda r: -r['amount'])

    # ------------------------------------------------------------------- the peer group
    # RULE 13 / rule 11's cousin: the proof carries NO denominator for any town, so the
    # LEVEL of free cash cannot be compared and is not. Composition is a share and does.
    peers = []
    for t in towns:
        ident = {y: at(t, y, IDENT) for y in years}
        peers.append(dict(
            town=t,
            identified=[dict(year=y, amount=round(ident[y], 2)) for y in years],
            certified=[dict(year=y, amount=round(at(t, y, CERT), 2)) for y in years],
            unspent_share=[dict(year=y, share=round(at(t, y, UNSPENT) / ident[y], 4))
                           for y in years],
            receipts_share=[dict(year=y, share=round(at(t, y, RECEIPTS) / ident[y], 4))
                            for y in years],
            aid_share=[dict(year=y, share=round(at(t, y, AID) / ident[y], 4))
                       for y in years],
            unspent_mean=round(statistics.mean(at(t, y, UNSPENT) / ident[y] for y in years), 4),
            receipts_positive=sum(1 for y in years if at(t, y, RECEIPTS) > 0),
            receipts_total=round(sum(at(t, y, RECEIPTS) for y in years), 2),
            aid_positive=sum(1 for y in years if at(t, y, AID) > 0),
            aid_abs_mean_share=round(
                statistics.mean(abs(at(t, y, AID)) / ident[y] for y in years), 4),
            aid_max_swing=round(max(abs(at(t, years[i + 1], AID) - at(t, years[i], AID))
                                    for i in range(len(years) - 1)), 2),
            certified_cv=round(
                statistics.pstdev([at(t, y, CERT) for y in years])
                / statistics.mean([at(t, y, CERT) for y in years]), 4),
        ))
    peers.sort(key=lambda p: -p['unspent_mean'])

    # 2025 against each town's OWN 2021-24 average. A ratio, so it compares; and it is
    # the measure that says whether the record year was an event or a stance.
    base_years = years[:-1]
    multiples = []
    for t in towns:
        avg = statistics.mean(at(t, y, UNSPENT) for y in base_years)
        multiples.append(dict(town=t, latest=round(at(t, years[-1], UNSPENT), 2),
                              average=round(avg, 2),
                              multiple=round(at(t, years[-1], UNSPENT) / avg, 2)))
    multiples.sort(key=lambda r: -r['multiple'])

    town_years = len(towns) * len(years)
    receipts_beat = sum(1 for t in towns for y in years if at(t, y, RECEIPTS) > 0)

    # ------------------------------------------- the model's typed constants, recomputed
    # Every one of these was a literal in model/freecash.py and none of them was checked
    # against the table it came out of. Rule 2 covers a constant in a module exactly as
    # it covers a sentence in a page.
    import freecash as fc  # noqa: E402
    for name, got, want in (
        ('CERTIFIED', at(TOWN, years[-1], CERT), fc.CERTIFIED),
        ('IDENTIFIED', at(TOWN, years[-1], IDENT), fc.IDENTIFIED),
        ('UNSPENT_2025', at(TOWN, years[-1], UNSPENT), fc.UNSPENT_2025),
        ('UNSPENT_AVG_2021_24',
         statistics.mean(at(TOWN, y, UNSPENT) for y in base_years),
         fc.UNSPENT_AVG_2021_24),
    ):
        if abs(got - want) > 0.5:
            fail(f'model/freecash.py states {name} = {want:,.2f}; the proof gives '
                 f'{got:,.2f}')
    typed = {t: m for t, m in fc.PEER_MULTIPLES}
    for r in multiples:
        if r['town'] not in typed:
            fail(f'model/freecash.py PEER_MULTIPLES has no {r["town"]}')
        if abs(typed[r['town']] - r['multiple']) > 0.005:
            fail(f'model/freecash.py types {r["town"]} at {typed[r["town"]]}; the proof '
                 f'gives {r["multiple"]}')

    # --------------------------------------------- CHECK 5: what the town actually said
    said = []
    for s in SAID:
        full = os.path.join(MINUTES, s['path'])
        if not os.path.exists(full):
            fail(f'{s["key"]}: {s["path"]} is not in the meeting archive')
        flat = ' '.join(open(full, encoding='utf-8', errors='replace').read().split())
        if ' '.join(s['quote'].split()) not in flat:
            fail(f'{s["key"]}: the quote is not in {s["path"]} — rule 15a puts what was '
                 'said beside the measurement, and a quote nobody can find is worse than '
                 'no quote')
        said.append(dict(key=s['key'], board=s['board'], date=s['date'],
                         agenda=s['agenda'], quote=s['quote'], beside=s['beside'],
                         cite=SITE + s['path'], town=BASE + s['suffix']))

    # --------------------------------- the Town's own proof, and where it lands instead
    rf = rollforward()
    versus = []
    for r in rf:
        y = r['fy']
        if y not in years:
            continue
        ident = at(TOWN, y, IDENT)
        cert = at(TOWN, y, CERT)
        versus.append(dict(
            year=y, as_of=r['as_of'],
            undesignated=r['total']['amount'], identified=round(ident, 2),
            certified=round(cert, 2),
            undesignated_minus_identified=round(r['total']['amount'] - ident, 2),
            identified_minus_certified=round(ident - cert, 2)))
    if not versus:
        fail('no year has both a Town roll-forward and a DLS proof — the join between the '
             'two publishers matched nothing, which reads exactly like the Town not '
             'publishing one')

    # ------------------------------------------------- CHECK 6: the gap register, by key
    reg = {(g['side'], g['what']): g['why']
           for g in (dict(r) for r in c.execute('SELECT side, what, why FROM money_gaps'))}
    gaps = []
    for side, what in GAP_KEYS:
        if (side, what) not in reg:
            fail(f'money_gaps has no row ({side}, {what!r}). This page quotes the register '
                 'by key so a renamed row fails here instead of rendering an empty box.')
        gaps.append(dict(side=side, what=what, why=reg[(side, what)]))
    # Rule 2 reaches the register. These two rows quote figures in their prose; the
    # figures are recomputed above and the row must still contain them, so a gap whose
    # numbers have drifted from the data fails here rather than being published as one.
    v25 = next(v for v in versus if v['year'] == years[-1])
    for what, values in (
        ('Why DLS certifies a different amount from the free cash its own proof identifies',
         [v25['identified_minus_certified']]),
        ('Why the Town’s undesignated fund balance and DLS’s free cash differ',
         [v25['undesignated'], v25['identified'], v25['certified']]),
    ):
        why = reg[('money_in', what)]
        for v in values:
            money = f'${v:,.2f}' if v % 1 else f'${v:,.0f}'
            if money not in why:
                fail(f'money_gaps ({what!r}) no longer states {money}, which is what the '
                     'data now gives. Rule 2 covers the register too.')

    return dict(
        generated_by='scripts/build_free_cash_charts.py',
        source='sources/data/lunenburg.db — free_cash_proof, money_gaps; the annual town '
               'reports for the Town’s own roll-forward; sources/meetings/text/ for what '
               'was said',
        document=dict(
            what='Free Cash Proof, Massachusetts Department of Revenue, Division of Local '
                 'Services',
            our_copy='/docs/state-dls/free-cash-proof-lunenburg.xlsx',
            publisher_filename='FCPCompareLunenburg.xlsx',
            provenance='/docs/state-dls/PROVENANCE.md',
            sheet='Sheet1', rows='4–17', year_columns='B3:F3',
            ref_note='free_cash_proof.source_ref names the LABEL cell in column A for '
                     'every year. The amount for a year sits in that year’s own column, '
                     'and it is that coordinate this page cites.'),
        coverage=dict(towns=towns, years=years, rows=len(raw),
                      town_years=town_years, lines=len(lines),
                      cells_checked_against_workbook=checked_cells),
        crossfoot=dict(
            tested=tied, tied=tied, ties_to=IDENT, against=CERT,
            certified_differs=len(differ),
            rows=crossfoot,
            chained=chained),
        proof=proof, line_order=lines,
        composition=composition, part_order=PART_ORDER,
        part_labels={key: label for _, key, label in NAMED},
        totals=dict(years=len(years), identified=round(total_ident, 2),
                    certified=round(sum(r['certified'] for r in composition), 2),
                    by_line=by_line),
        peers=peers,
        multiples=multiples,
        receipts=dict(town_years=town_years, beat_estimate=receipts_beat,
                      towns=len(towns), years=len(years)),
        denominator_note=fc.PEER_DENOMINATOR_NOTE,
        rollforward=rf,
        versus_dls=versus,
        said=said,
        gaps=gaps,
    )


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--check', action='store_true',
                    help='fail if the published file is not what this would write')
    a = ap.parse_args()
    payload = json.dumps(build(), indent=1, sort_keys=True) + '\n'
    rel = os.path.relpath(OUT, ROOT)
    if a.check:
        have = open(OUT, encoding='utf-8').read() if os.path.exists(OUT) else ''
        if have != payload:
            print(f'STALE — {rel} is not what the database now produces. '
                  'Run scripts/build_free_cash_charts.py.')
            return 1
        print(f'ok — {rel} reproduces from the database')
        return 0
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    open(OUT, 'w', encoding='utf-8').write(payload)
    d = json.loads(payload)
    print(f'wrote {rel} — {d["crossfoot"]["tied"]} of {d["crossfoot"]["tested"]} town-years '
          f'cross-foot to “{d["crossfoot"]["ties_to"]}”, '
          f'{d["crossfoot"]["certified_differs"]} of them certify a different amount; '
          f'{d["receipts"]["beat_estimate"]} of {d["receipts"]["town_years"]} town-years '
          f'beat their local receipt estimate; '
          f'{d["coverage"]["cells_checked_against_workbook"]} cells checked against the '
          f'workbook')
    return 0


if __name__ == '__main__':
    sys.exit(main())
