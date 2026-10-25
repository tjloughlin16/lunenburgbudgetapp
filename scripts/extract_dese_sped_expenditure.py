#!/usr/bin/env python3
"""Lunenburg's special education spending as DESE prints it, FY2008-FY2025, tied to itself.

    python3 scripts/extract_dese_sped_expenditure.py           # write the CSV
    python3 scripts/extract_dese_sped_expenditure.py --check   # fail if it no longer reproduces

Reads `sources/state-dese/radar-sped-direct-expenditure-trends-fy2008-fy2025.xlsx` (DESE
RADAR, *Special Education Direct Expenditure Trends*; landed by
fetch_dese_sped_expenditure.py) and writes `sources/data/dese-sped-expenditure.csv`: one
row per fiscal year for Lunenburg (LEA 0162) and one for DESE's State Total row.

THE COLUMNS ARE DESE'S, IN DESE'S GROUPS

  in district      teaching, other_instructional, in_district_transportation
  out of district  mass_public_and_collaboratives, private_and_out_of_state,
                   out_of_district_transportation
  other            nonpublic_health, grants_revolving
  totals           sped_total (DESE's own SUM of the eight), total_expenditures,
                   sped_share (DESE's own M/N)

DESE's statewide report to the Legislature (research/reports/2020/09-sped-fy2019.docx)
defines the groups: the first seven are GENERAL FUND spending, *including school committee
and city or town appropriations, and expenditures from the Special Education Reimbursement
(Circuit Breaker) Program*; `grants_revolving` is special education spending from grants
such as IDEA and from revolving accounts, and is NOT split in district / out of district.

A BLANK IS NOT A ZERO. Where the workbook's cell is empty the CSV is empty. DESE's own SUM
treats an empty cell as zero, and so does the sped_total check below, but a reader should
not: FY2024 in-district transportation is empty for Lunenburg while every other year
carries six figures, and that is a question for DESE, not a measurement of $0.

WHAT EVERY ROW IS TIED TO, and the script refuses to write if any tie fails (rule 13):

  1. The eight categories sum to the workbook's own cached `Combined Special Ed
     Expenditures` (column M, a SUM formula DESE computed).
  2. The workbook's own cached share (column O) equals M / N.
  3. FY2008-FY2023: every category, and total expenditures, equals the same figure in the
     hidden long-format sheet `Spec_Education_Exp run 03192024` -- DESE's underlying run,
     keyed by category code, with `TOTAL OPERATING` as code B. A second printing of the
     same numbers in a different shape. FY2024-FY2025 are not in that run.
  4. FY2025: every figure equals the `FY25 comparison` sheet's cached row.
  5. The State Total row equals the sum of every district row for that year, per column.

WHAT IT DOES NOT GIVE: any count of children. There is no per-pupil figure in this
workbook. A per-child figure is a derivation using another file's denominator and must be
labelled ours.
"""
import argparse
import csv
import io
import os
import sys
from collections import defaultdict

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOC = 'sources/state-dese/radar-sped-direct-expenditure-trends-fy2008-fy2025.xlsx'
OUT = os.path.join(ROOT, 'sources', 'data', 'dese-sped-expenditure.csv')
LEA = '0162'
STATE = 999

# data-sheet column letter -> (csv column, code in the hidden long sheet)
CATS = [('E', 'teaching', '1'),
        ('F', 'other_instructional', '2'),
        ('G', 'in_district_transportation', '3'),
        ('H', 'mass_public_and_collaboratives', '4'),
        ('I', 'private_and_out_of_state', '5'),
        ('J', 'out_of_district_transportation', '6'),
        ('K', 'nonpublic_health', '7'),
        ('L', 'grants_revolving', '8')]
# The header the columns are read under. Asserted, never inferred (rule 13b.4).
HEADER = {'B': 'LEA', 'C': 'District Name', 'D': 'Fiscal Year', 'E': 'Teaching',
          'F': 'Other Instructional', 'G': 'In-District Transportation',
          'H': 'Mass. Public Schools and Collaboratives',
          'I': 'Mass Private and Out-of-State Schools',
          'J': 'Out-of-District Transportation', 'K': 'Non-public Health Services',
          'L': 'Spending from Grants and Revolving Funds',
          'M': 'Combined Special Ed Expenditures (sum of E through L)',
          'N': 'Total Expenditures',
          'O': 'Special Education Percentage of Total Expenditures (M as % of N)'}
HEADER_ROW = 6
FIELDS = (['fy', 'lea', 'district', 'level'] + [c for _, c, _ in CATS] +
          ['sped_total', 'total_expenditures', 'sped_share', 'reconciles', 'doc_id'])
TOL = 0.01


def col(letter):
    return ord(letter) - ord('A')


def num(v):
    if v is None or v == '':
        return None
    if isinstance(v, str):
        v = v.strip()
        if v == '':
            return None
        return float(v)
    return float(v)


def fmt(v):
    if v is None:
        return ''
    r = round(v, 6)
    return ('%d' % r) if r == int(r) else repr(r)


def read():
    import openpyxl
    path = os.path.join(ROOT, DOC)
    if not os.path.exists(path):
        raise SystemExit('%s is not on disk. python3 scripts/sync_archive.py --pull' % DOC)
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)

    rows = list(wb['data'].iter_rows(values_only=True))
    head = rows[HEADER_ROW - 1]
    for letter, want in HEADER.items():
        if head[col(letter)] != want:
            raise SystemExit('data!%s%d is %r, expected %r. The layout moved; read the '
                             'header again before extracting.'
                             % (letter, HEADER_ROW, head[col(letter)], want))
    data = [r for r in rows[HEADER_ROW:] if r[col('D')] is not None]

    long = defaultdict(dict)    # (lea4, fy) -> {code: value}
    for r in wb['Spec_Education_Exp run 03192024'].iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        long[(str(r[1]), int(r[0]))][str(r[3])] = r[5]

    fy25 = {}
    for r in wb['FY25 comparison'].iter_rows(min_row=8, values_only=True):
        if isinstance(r[1], int):
            fy25[r[1]] = r
    wb.close()
    return data, long, fy25


def build():
    data, long, fy25 = read()
    fails = []
    by_year_sum = defaultdict(lambda: defaultdict(float))
    picked = []
    for r in data:
        lea_raw, fy = r[col('B')], int(r[col('D')])
        is_state = lea_raw == STATE
        if not is_state:
            for letter in [c[0] for c in CATS] + ['N']:
                by_year_sum[fy][letter] += num(r[col(letter)]) or 0.0
        if is_state or str(lea_raw) == LEA:
            picked.append(r)

    out = []
    for r in picked:
        lea_raw, fy = r[col('B')], int(r[col('D')])
        is_state = lea_raw == STATE
        where = '%s FY%d' % ('State Total' if is_state else 'Lunenburg', fy)
        vals = {c: num(r[col(L)]) for L, c, _ in CATS}
        m, n, o = num(r[col('M')]), num(r[col('N')]), num(r[col('O')])
        if m is None or n is None:
            fails.append('%s: no cached total (M=%r, N=%r)' % (where, m, n))
            continue
        # 1. DESE's own SUM
        s = sum(v or 0.0 for v in vals.values())
        if abs(s - m) > TOL:
            fails.append('%s: categories sum to %.2f, the workbook prints %.2f' % (where, s, m))
        # 2. DESE's own share
        if o is None or abs(o - m / n) > 1e-9:
            fails.append('%s: share %r is not M/N %r' % (where, o, m / n))
        # 3. the hidden run, FY2008-FY2023
        if not is_state:
            name = r[col('C')]
            if name not in ('Lunenburg', '', None):
                fails.append('%s: district printed as %r' % (where, name))
            run = long.get((LEA, fy))
            if fy <= 2023:
                if not run:
                    fails.append('%s: absent from the hidden run' % where)
                else:
                    for L, c, code in CATS:
                        a, b = vals[c], num(run.get(code))
                        if abs((a or 0.0) - (b or 0.0)) > TOL:
                            fails.append('%s %s: data sheet %r, run %r' % (where, c, a, b))
                    if abs(n - num(run.get('B'))) > TOL:
                        fails.append('%s: total %r, run TOTAL OPERATING %r'
                                     % (where, n, run.get('B')))
            # 4. the FY25 comparison sheet
            if fy == 2025:
                c25 = fy25.get(int(LEA))
                if not c25:
                    fails.append('%s: absent from FY25 comparison' % where)
                else:
                    # FY25 comparison columns D..N carry E..O of the data sheet.
                    for i, L in enumerate('EFGHIJKLMNO'):
                        a = num(r[col(L)]) or 0.0
                        b = num(c25[3 + i]) or 0.0
                        if abs(a - b) > TOL:
                            fails.append('%s col %s: data %r, FY25 comparison %r'
                                         % (where, L, a, b))
        else:
            # 5. the State Total is the sum of the districts
            for L in [c[0] for c in CATS] + ['N']:
                a, b = num(r[col(L)]) or 0.0, by_year_sum[fy][L]
                if abs(a - b) > max(1.0, abs(a) * 1e-9):
                    fails.append('%s col %s: printed %.2f, districts sum to %.2f'
                                 % (where, L, a, b))
        row = dict(fy=fy, lea='00000000' if is_state else '01620000',
                   district='State' if is_state else 'Lunenburg',
                   level='state' if is_state else 'district',
                   sped_total=fmt(m), total_expenditures=fmt(n), sped_share=fmt(o),
                   reconciles='yes', doc_id=DOC)
        row.update({c: fmt(v) for c, v in vals.items()})
        out.append(row)
    out.sort(key=lambda d: (d['level'] != 'state', d['fy']))
    return out, fails


def render(rows):
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=FIELDS, lineterminator='\n')
    w.writeheader()
    w.writerows(rows)
    return buf.getvalue()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--check', action='store_true')
    a = ap.parse_args()
    rows, fails = build()
    if fails:
        for f in fails:
            print('  !!', f)
        print('%d ties failed. Nothing written.' % len(fails))
        return 1
    years = sorted({r['fy'] for r in rows if r['level'] == 'district'})
    text = render(rows)
    if a.check:
        have = open(OUT, encoding='utf-8', newline='').read() if os.path.exists(OUT) else ''
        if have != text:
            print('STALE -- %s no longer reproduces' % os.path.relpath(OUT, ROOT))
            return 1
        print('ok -- %d rows, Lunenburg FY%d-FY%d, every row tied to the workbook\'s own '
              'totals' % (len(rows), years[0], years[-1]))
        return 0
    tmp = OUT + '.tmp'
    with open(tmp, 'w', encoding='utf-8', newline='') as fh:
        fh.write(text)
    os.replace(tmp, OUT)
    print('wrote %s: %d rows, Lunenburg FY%d-FY%d, every row tied'
          % (os.path.relpath(OUT, ROOT), len(rows), years[0], years[-1]))
    return 0


if __name__ == '__main__':
    sys.exit(main())
